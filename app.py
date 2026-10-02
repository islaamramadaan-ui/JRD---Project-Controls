
import streamlit as st
import pandas as pd
import numpy as np
from datetime import datetime, timedelta, date
from pathlib import Path
from io import BytesIO
import base64, math, os
from openpyxl import load_workbook

import plotly.express as px
import plotly.graph_objects as go
import matplotlib.pyplot as plt
import matplotlib.dates as mdates

# PDF/PPT
from reportlab.lib.pagesizes import A4, A3, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

st.set_page_config(
    page_title="JRD Project Controls Hub",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ------------------------------------------------------------------
# Brand
# ------------------------------------------------------------------
BRAND_GREEN = "#07795C"
BRAND_SILVER = "#D1D3D4"
BRAND_BLACK = "#121212"
BRAND_YELLOW = "#FFC400"
BRAND_OFFWHITE = "#F7F7F5"
BRAND_RED = "#D62828"

ROOT = Path(__file__).parent if "__file__" in globals() else Path(".")
ASSETS = ROOT / "assets"
ALTA_LOGO = ASSETS / "alta_logo.png"
DDDC_LOGO = ASSETS / "dddc_logo.png"
DAR_LOGO = ASSETS / "dar_logo.png"
COVER_IMG = ASSETS / "jrd_cover.png"
DB_TEMPLATE = ROOT / "templates" / "JRD_Project_Data_Input_v1.xlsx"

st.markdown(f"""
<style>
html, body, [class*="css"] {{ font-family: Poppins, Arial, sans-serif; }}
.block-container {{padding-top: 1.15rem; padding-bottom: 3rem;}}
[data-testid="stSidebar"] {{background: {BRAND_BLACK};}}
[data-testid="stSidebar"] * {{color: {BRAND_OFFWHITE};}}
div[data-testid="stMetric"] {{
  border: 1px solid rgba(209,211,212,.35);
  border-left: 5px solid {BRAND_GREEN};
  padding: 12px 14px; border-radius: 10px;
  background: rgba(247,247,245,.04);
}}
.kpi-title {{font-size:.74rem; opacity:.72; text-transform:uppercase; letter-spacing:.04em;}}
.kpi-big {{font-size:1.65rem; font-weight:700;}}
.section-title {{
  margin-top: 10px; padding: 8px 12px; border-radius: 6px;
  background:{BRAND_GREEN}; color:white; font-weight:700;
}}
.neg {{color:{BRAND_RED}; font-weight:700;}}
.pos {{color:{BRAND_GREEN}; font-weight:700;}}
.warn {{color:{BRAND_YELLOW}; font-weight:700;}}
.cover-card {{
  position:relative; border-radius:16px; overflow:hidden; margin-bottom:14px;
  border:1px solid rgba(209,211,212,.25);
}}
.cover-overlay {{
  position:absolute; inset:0; background:linear-gradient(90deg,rgba(18,18,18,.88),rgba(18,18,18,.22));
  padding:38px; display:flex; flex-direction:column; justify-content:flex-end;
}}
.cover-title {{font-size:2.4rem; font-weight:800; color:white; line-height:1.0;}}
.cover-sub {{color:{BRAND_SILVER}; font-size:1rem; margin-top:8px;}}
.small {{font-size:.82rem; opacity:.72;}}
</style>
""", unsafe_allow_html=True)

# ------------------------------------------------------------------
# Helpers
# ------------------------------------------------------------------
def img_b64(path):
    if not path.exists():
        return ""
    return base64.b64encode(path.read_bytes()).decode()

def fmt_date(x):
    if pd.isna(x): return "—"
    return pd.Timestamp(x).strftime("%d-%b-%Y")

def to_dt(s): return pd.to_datetime(s, errors="coerce")
def to_num(s): return pd.to_numeric(s, errors="coerce")

def load_schedule(uploaded):
    xls = pd.ExcelFile(uploaded)
    if "TASK" not in xls.sheet_names or "TASKPRED" not in xls.sheet_names:
        raise ValueError("Workbook must contain TASK and TASKPRED sheets.")
    task = pd.read_excel(xls, sheet_name="TASK")
    pred = pd.read_excel(xls, sheet_name="TASKPRED")
    if "task_code" not in task.columns or "pred_task_id" not in pred.columns:
        raise ValueError("Expected Primavera technical column names were not found.")

    task = task[task["task_code"].astype(str).str.strip() != "Activity ID"].copy()
    pred = pred[pred["pred_task_id"].astype(str).str.strip() != "Predecessor"].copy()

    for c in ["task_code","task_name","status_code","wbs_id"]:
        if c in task.columns:
            task[c] = task[c].astype(str).str.strip()

    task["duration_days"] = to_num(task.get("target_drtn_hr_cnt"))
    task["total_float_days"] = to_num(task.get("total_float_hr_cnt"))
    task["start"] = to_dt(task.get("start_date"))
    task["finish"] = to_dt(task.get("end_date"))

    aliases = {
        "actual_start":["act_start_date","actual_start_date"],
        "actual_finish":["act_end_date","actual_finish_date"],
        "remain_duration":["remain_drtn_hr_cnt","remaining_duration"],
        "physical_pct":["phys_complete_pct","physical_complete_pct"],
        "pct_complete":["complete_pct","task_complete_pct"],
        "budget_cost":["target_cost","budgeted_cost","planned_cost","budget_cost"],
        "actual_cost":["act_cost","actual_cost"],
        "budget_units":["target_qty","budgeted_units","planned_units","budget_units","target_labor_units"],
        "actual_units":["act_qty","actual_units","actual_labor_units"],
        "remain_units":["remain_qty","remaining_units","remain_labor_units"],
    }
    for out, candidates in aliases.items():
        task[out] = np.nan
        for c in candidates:
            if c in task.columns:
                task[out] = task[c]; break
    for c in ["actual_start","actual_finish"]:
        task[c] = to_dt(task[c])
    for c in ["remain_duration","physical_pct","pct_complete","budget_cost","actual_cost","budget_units","actual_units","remain_units"]:
        task[c] = to_num(task[c])

    pred["pred_task_id"] = pred["pred_task_id"].astype(str).str.strip()
    pred["task_id"] = pred["task_id"].astype(str).str.strip()
    pred["pred_type"] = pred["pred_type"].astype(str).str.strip()
    pred["lag_days"] = to_num(pred.get("lag_hr_cnt"))

    optional = {}
    for s in xls.sheet_names:
        su = s.upper().replace("_"," ").strip()
        if su in {"TASKRSRC","RESOURCEASSIGNMENT","RESOURCE ASSIGNMENT","RESASSIGN","ASSIGNMENTS"}:
            optional["assignments"] = pd.read_excel(xls, sheet_name=s)
        elif su in {"RSRC","RESOURCES","RESOURCE"}:
            optional["resources"] = pd.read_excel(xls, sheet_name=s)
    return task, pred, optional

def category_from_row(row):
    code = str(row.get("task_code","")).upper()
    name = str(row.get("task_name","")).lower()
    wbs = str(row.get("wbs_id","")).lower()
    if code.startswith("PQ-") or "prequalification" in name: return "Prequalification"
    if code.startswith("SD-") or "shop drawing" in name: return "Shop Drawings"
    if code.startswith("MS-") or "material submittal" in name or "material approval" in name: return "Material Submittals"
    if code.startswith(("PR-","PO-","MF-","DL-")) or any(k in name for k in ["procurement","purchase order","lpo","manufacturing","fabrication","delivery on site"]): return "Procurement"
    if code.startswith("AUTH-") or "authority" in name or " noc" in f" {name}": return "Authorities"
    if code.startswith("TC-") or any(k in name for k in ["testing","commissioning"]): return "T&C"
    if code.startswith("HO-") or any(k in name for k in ["handover","taking over","snag","o&m","as-built","close-out","closeout"]): return "Handover"
    if any(k in name for k in ["excavation","raft","pile","slab","column","wall","blockwork","plaster","screed","ceiling","tile","paint","façade","facade","external works"]): return "Construction"
    if any(k in wbs for k in ["construction","substructure","superstructure","finishes","external"]): return "Construction"
    return "Other"

def package_name(row):
    name = str(row.get("task_name",""))
    # remove common prefixes / action words to group major packages
    cleaned = name
    for token in ["Prequalification","Shop Drawing","Shop Drawings","Material Submittal","Material Approval",
                  "Submit","Submission","Approval","Procurement","Purchase Order","LPO","Manufacturing","Fabrication",
                  "Delivery on Site","Delivery"]:
        cleaned = cleaned.replace(token,"")
    cleaned = cleaned.replace(" - "," ").replace(":"," ").strip(" -")
    return cleaned[:90] if cleaned else name[:90]

def enrich(task, pred):
    t = task.copy()
    t["category"] = t.apply(category_from_row, axis=1)
    t["package"] = t.apply(package_name, axis=1)
    pc = pred.groupby("task_id").size()
    sc = pred.groupby("pred_task_id").size()
    t["has_pred"] = t["task_code"].isin(pc.index)
    t["has_succ"] = t["task_code"].isin(sc.index)
    return t

def project_facts(task):
    return {
        "start": task["start"].min(),
        "finish": task["finish"].max(),
        "activities": len(task),
        "critical": int((task["total_float_days"] <= 0).sum()),
        "negative": int((task["total_float_days"] < 0).sum()),
    }

def planned_curve(task, value_col, freq="W"):
    data = task.dropna(subset=["start","finish",value_col]).copy()
    data = data[data[value_col].fillna(0) != 0]
    if data.empty: return pd.DataFrame()
    parts=[]
    for _,r in data.iterrows():
        dates=pd.date_range(r["start"].normalize(),r["finish"].normalize(),freq="D")
        if len(dates)==0: continue
        daily=float(r[value_col])/len(dates)
        parts.append(pd.DataFrame({"date":dates,"planned":daily}))
    if not parts: return pd.DataFrame()
    out=pd.concat(parts).groupby("date",as_index=False)["planned"].sum()
    out=out.set_index("date").resample(freq)["planned"].sum().reset_index()
    out["cum_planned"]=out["planned"].cumsum()
    total=out["planned"].sum()
    out["cum_planned_pct"]=np.where(total, out["cum_planned"]/total*100,0)
    return out

def assignment_to_activity(assignments):
    if assignments is None or assignments.empty:
        return pd.DataFrame()
    cols={str(c).lower():c for c in assignments.columns}
    def pick(options):
        for o in options:
            if o in cols: return cols[o]
        return None
    act=pick(["task_code","activity id","activity_id","task_id"])
    s=pick(["start_date","start","task_start_date"])
    f=pick(["end_date","finish","finish_date","task_finish_date"])
    pc=pick(["target_cost","budgeted_cost","planned_cost","budget_cost"])
    ac=pick(["act_cost","actual_cost"])
    pu=pick(["target_qty","budgeted_units","planned_units","budget_units","target_labor_units"])
    au=pick(["act_qty","actual_units","actual_labor_units"])
    rn=pick(["resource_name","resource name","rsrc_name"])
    rid=pick(["resource_id","resource id","rsrc_id"])
    out=pd.DataFrame()
    if act: out["activity_id"]=assignments[act].astype(str).str.strip()
    if s: out["start"]=to_dt(assignments[s])
    if f: out["finish"]=to_dt(assignments[f])
    if pc: out["planned_cost"]=to_num(assignments[pc])
    if ac: out["actual_cost"]=to_num(assignments[ac])
    if pu: out["planned_units"]=to_num(assignments[pu])
    if au: out["actual_units"]=to_num(assignments[au])
    if rn: out["resource_name"]=assignments[rn].astype(str)
    if rid: out["resource_id"]=assignments[rid].astype(str)
    return out

def timephase_assign(assign_df, val_col, data_date, freq="W"):
    if assign_df is None or assign_df.empty or val_col not in assign_df.columns:
        return pd.DataFrame()
    d=assign_df.dropna(subset=["start","finish",val_col]).copy()
    d=d[d[val_col].fillna(0)!=0]
    if d.empty:return pd.DataFrame()
    rows=[]
    for _,r in d.iterrows():
        dates=pd.date_range(r["start"].normalize(),r["finish"].normalize(),freq="D")
        if len(dates)==0:continue
        daily=float(r[val_col])/len(dates)
        tmp=pd.DataFrame({"date":dates,val_col:daily})
        rows.append(tmp)
    out=pd.concat(rows).groupby("date",as_index=False)[val_col].sum()
    out=out.set_index("date").resample(freq)[val_col].sum().reset_index()
    return out

def merge_plan_actual(plan, actual, freq, planned_name="planned", actual_name="actual"):
    # Robust handling when only Planned or only Actual data exists.
    # Always preserve a date axis and create the missing series as zero.
    if (plan is None or plan.empty) and (actual is None or actual.empty):
        return pd.DataFrame(columns=["date", planned_name, actual_name,
                                     "cum_planned", "cum_actual",
                                     "cum_planned_pct", "cum_actual_pct"])

    if plan is None or plan.empty:
        out = actual.copy()
        if "date" not in out.columns:
            return pd.DataFrame()
        if actual_name not in out.columns:
            value_cols = [c for c in out.columns if c != "date"]
            if value_cols:
                out = out.rename(columns={value_cols[0]: actual_name})
        out[planned_name] = 0.0

    elif actual is None or actual.empty:
        out = plan.copy()
        if "date" not in out.columns:
            return pd.DataFrame()
        if planned_name not in out.columns:
            value_cols = [c for c in out.columns if c != "date"]
            if value_cols:
                out = out.rename(columns={value_cols[0]: planned_name})
        out[actual_name] = 0.0

    else:
        out = pd.merge(plan, actual, on="date", how="outer")

    out["date"] = pd.to_datetime(out["date"], errors="coerce")
    out = out.dropna(subset=["date"]).sort_values("date").fillna(0)

    if planned_name not in out.columns:
        out[planned_name] = 0.0
    if actual_name not in out.columns:
        out[actual_name] = 0.0

    out[planned_name] = pd.to_numeric(out[planned_name], errors="coerce").fillna(0.0)
    out[actual_name] = pd.to_numeric(out[actual_name], errors="coerce").fillna(0.0)

    out["cum_planned"] = out[planned_name].cumsum()
    out["cum_actual"] = out[actual_name].cumsum()

    pt = out[planned_name].sum()
    out["cum_planned_pct"] = np.where(pt != 0, out["cum_planned"] / pt * 100, 0.0)
    # Actual progress percentage is measured against the total planned basis.
    out["cum_actual_pct"] = np.where(pt != 0, out["cum_actual"] / pt * 100, 0.0)

    return out.reset_index(drop=True)


def read_any_excel(uploaded):
    """Return all sheets as raw dataframes. Handles Streamlit UploadedFile or local path."""
    if uploaded is None:
        return {}
    try:
        xls = pd.ExcelFile(uploaded)
        return {s: pd.read_excel(xls, sheet_name=s, header=None) for s in xls.sheet_names}
    except Exception:
        return {}

def _excel_serial_date(x):
    if pd.isna(x): return pd.NaT
    if isinstance(x,(pd.Timestamp,datetime,date)): return pd.Timestamp(x)
    try:
        if isinstance(x,(int,float,np.integer,np.floating)) and 20000 <= float(x) <= 80000:
            return pd.Timestamp("1899-12-30") + pd.to_timedelta(float(x), unit="D")
    except Exception: pass
    return pd.to_datetime(x, errors="coerce")

def _mixed_dates(values):
    return pd.Series([_excel_serial_date(v) for v in list(values)])

def _is_date_like(x):
    y=_excel_serial_date(x)
    return pd.notna(y) and 1990 <= y.year <= 2050

def parse_planned_curve_excel(uploaded, preferred=("PLANNED", "W"), default_freq="W"):
    """
    Parses common P6 planned progress Excel curves:
    - Sheets like PLANNED PROGRESS W/M where a date row is followed by a values row.
    - Sheets like WEEKLY/MONTHLY where row contains dates and following row contains spread values.
    Output columns: date, planned, cum_planned, cum_planned_pct
    """
    sheets = read_any_excel(uploaded)
    if not sheets:
        return pd.DataFrame(columns=["date","planned","cum_planned","cum_planned_pct"])

    scored = []
    pref = [p.upper() for p in preferred]
    for sname, df in sheets.items():
        sn = sname.upper()
        sheet_bonus = sum(10 for p in pref if p in sn)
        if df.empty:
            continue
        for r in range(min(len(df), 60)):
            row = df.iloc[r]
            date_cols = [c for c, v in row.items() if _is_date_like(v)]
            if len(date_cols) < 3:
                continue
            for offset in range(1, 6):
                if r + offset >= len(df):
                    continue
                vrow = df.iloc[r + offset]
                vals = pd.to_numeric(vrow[date_cols], errors="coerce")
                numeric_count = vals.notna().sum()
                positive_sum = vals.fillna(0).abs().sum()
                if numeric_count >= max(3, len(date_cols) * 0.45) and positive_sum > 0:
                    score = sheet_bonus + len(date_cols) + numeric_count - offset
                    scored.append((score, sname, r, r + offset, date_cols))
    if not scored:
        return pd.DataFrame(columns=["date","planned","cum_planned","cum_planned_pct"])

    scored.sort(reverse=True, key=lambda x: x[0])
    _, sname, date_row, value_row, date_cols = scored[0]
    df = sheets[sname]
    dates = _mixed_dates(df.loc[date_row, date_cols].values)
    values = pd.to_numeric(df.loc[value_row, date_cols], errors="coerce").fillna(0.0)
    out = pd.DataFrame({"date": dates.values, "planned": values.values})
    out = out.dropna(subset=["date"]).sort_values("date")
    out = out.groupby("date", as_index=False)["planned"].sum()
    out["cum_planned"] = out["planned"].cumsum()
    total = out["planned"].sum()
    out["cum_planned_pct"] = np.where(total != 0, out["cum_planned"] / total * 100, 0.0)
    return out.reset_index(drop=True)

def parse_manpower_plan_excel(uploaded, monthly=False):
    sheets = read_any_excel(uploaded)
    if not sheets:
        return pd.DataFrame(columns=["date","planned"])
    target_words = ["MONTHLY MPH", "MONTHLY"] if monthly else ["WEEKLY MPH", "WEEKLY"]
    candidates=[]
    for sname, df in sheets.items():
        sn=sname.upper()
        if not any(w in sn for w in target_words):
            continue
        for r in range(min(len(df), 25)):
            label = str(df.iloc[r,0]).upper() if df.shape[1] else ""
            if "PLANNED MANPOWER" in label:
                date_row = 0
                date_cols=[c for c,v in df.iloc[date_row].items() if _is_date_like(v)]
                if len(date_cols) >= 3:
                    vals=pd.to_numeric(df.loc[r,date_cols],errors="coerce").fillna(0.0)
                    out=pd.DataFrame({"date":_mixed_dates(df.loc[date_row,date_cols].values).values,
                                      "planned":vals.values})
                    out=out.dropna(subset=["date"]).sort_values("date")
                    candidates.append((len(date_cols), out))
    if candidates:
        candidates.sort(reverse=True, key=lambda x:x[0])
        return candidates[0][1].reset_index(drop=True)
    pc=parse_planned_curve_excel(uploaded, preferred=("MPH","WEEKLY" if not monthly else "MONTHLY"))
    return pc[["date","planned"]] if not pc.empty else pd.DataFrame(columns=["date","planned"])

def actual_curve_from_update(task, value_col, data_date, freq="W"):
    if task is None or task.empty or value_col not in task.columns:
        return pd.DataFrame(columns=["date","actual"])
    d=task.dropna(subset=[value_col]).copy()
    d=d[d[value_col].fillna(0)!=0]
    if d.empty or "actual_start" not in d.columns:
        return pd.DataFrame(columns=["date","actual"])
    rows=[]
    for _,r in d.iterrows():
        s=r.get("actual_start", pd.NaT)
        f=r.get("actual_finish", pd.NaT)
        if pd.isna(s):
            continue
        if pd.isna(f):
            f=data_date
        dates=pd.date_range(pd.Timestamp(s).normalize(), pd.Timestamp(f).normalize(), freq="D")
        if len(dates)==0:
            continue
        daily=float(r[value_col])/len(dates)
        rows.append(pd.DataFrame({"date":dates,"actual":daily}))
    if not rows:
        return pd.DataFrame(columns=["date","actual"])
    out=pd.concat(rows).groupby("date",as_index=False)["actual"].sum()
    return out.set_index("date").resample(freq)["actual"].sum().reset_index()

def combine_planned_actual(planned_curve, actual_curve):
    pc = planned_curve.copy() if planned_curve is not None else pd.DataFrame()
    ac = actual_curve.copy() if actual_curve is not None else pd.DataFrame()
    if not pc.empty:
        pc = pc[["date","planned"]].copy()
    if not ac.empty:
        ac = ac[["date","actual"]].copy()
    return merge_plan_actual(pc, ac, "X", "planned", "actual")

def make_download_chart(fig, width=1200, height=620):
    try:
        return fig.to_image(format="png", width=width, height=height, scale=2)
    except Exception:
        return None


def read_project_database(uploaded):
    """Read the user-maintained project database workbook."""
    if uploaded is None:
        return {}, {}
    try:
        xls = pd.ExcelFile(uploaded)
        sheets = {name: pd.read_excel(xls, sheet_name=name) for name in xls.sheet_names}
        setup = {}
        if "Project Setup" in sheets:
            df = sheets["Project Setup"]
            # locate Field/Value columns even when title rows exist
            if "Field" in df.columns and "Value" in df.columns:
                use = df[["Field","Value"]].dropna(subset=["Field"])
            else:
                raw = pd.read_excel(uploaded, sheet_name="Project Setup", header=None)
                hdr_idx = None
                for i in range(min(10,len(raw))):
                    row = raw.iloc[i].astype(str).str.strip().str.lower().tolist()
                    if "field" in row and "value" in row:
                        hdr_idx=i; break
                if hdr_idx is not None:
                    df2=pd.read_excel(uploaded, sheet_name="Project Setup", header=hdr_idx)
                    use=df2[["Field","Value"]].dropna(subset=["Field"])
                else:
                    use=pd.DataFrame(columns=["Field","Value"])
            for _,r in use.iterrows():
                setup[str(r["Field"]).strip()] = r["Value"]
        return setup, sheets
    except Exception:
        return {}, {}

def db_sheet(db_sheets, name, cols=None):
    df=db_sheets.get(name, pd.DataFrame()).copy()
    if df.empty:
        return df
    df=df.dropna(how="all")
    if cols:
        keep=[c for c in cols if c in df.columns]
        return df[keep].copy() if keep else pd.DataFrame()
    return df

def get_setup(setup, key, default=None):
    v=setup.get(key, default)
    return default if pd.isna(v) else v

def build_history_curves(db_sheets, data_date, current_mhr_actual=np.nan, current_cost_actual=np.nan,
                         current_mhr_plan=np.nan, current_cost_plan=np.nan):
    hist=db_sheet(db_sheets,"Update History")
    cols=["Data Date","Report No.","MHR Cumulative Plan %","MHR Cumulative Actual %",
          "Cost Cumulative Plan %","Cost Cumulative Actual %","Forecast Completion",
          "Negative Float Activities","Notes"]
    if hist.empty:
        hist=pd.DataFrame(columns=cols)
    for c in cols:
        if c not in hist.columns: hist[c]=np.nan
    hist=hist[cols].copy()
    hist["Data Date"]=pd.to_datetime(hist["Data Date"],errors="coerce")
    hist=hist.dropna(subset=["Data Date"])
    current={
        "Data Date":pd.Timestamp(data_date),
        "MHR Cumulative Plan %":current_mhr_plan,
        "MHR Cumulative Actual %":current_mhr_actual,
        "Cost Cumulative Plan %":current_cost_plan,
        "Cost Cumulative Actual %":current_cost_actual,
    }
    # replace same data date, otherwise append
    hist=hist[hist["Data Date"].dt.normalize()!=pd.Timestamp(data_date).normalize()]
    hist=pd.concat([hist,pd.DataFrame([current])],ignore_index=True).sort_values("Data Date")
    for c in ["MHR Cumulative Plan %","MHR Cumulative Actual %","Cost Cumulative Plan %","Cost Cumulative Actual %"]:
        hist[c]=pd.to_numeric(hist[c],errors="coerce")
    for prefix in ["MHR","Cost"]:
        ac=f"{prefix} Cumulative Actual %"; pc=f"{prefix} Cumulative Plan %"
        hist[f"{prefix} Weekly Actual %"]=hist[ac].diff().fillna(hist[ac])
        hist[f"{prefix} Weekly Plan %"]=hist[pc].diff().fillna(hist[pc])
    return hist.reset_index(drop=True)

def dataframe_or_message(df, message="Data not provided"):
    return df if df is not None and not df.empty else pd.DataFrame({"Status":[message]})

def style_float(v):
    if pd.isna(v): return ""
    return f"color: {BRAND_RED}; font-weight:700;" if v < 0 else ""

def date_window(task, dd, weeks):
    end=dd+pd.Timedelta(weeks=weeks)
    return task[
        ((task["start"]>=dd)&(task["start"]<=end)) |
        ((task["finish"]>=dd)&(task["finish"]<=end)) |
        ((task["start"]<=dd)&(task["finish"]>=dd))
    ].sort_values(["start","finish"])

def status_from_float(tf):
    if pd.isna(tf): return "Unknown"
    if tf < 0: return "DELAY / NEGATIVE FLOAT"
    if tf <= 7: return "Critical"
    if tf <= 20: return "High Risk"
    return "Normal"

def procurement_action(row, dd):
    tf=row.get("total_float_days",np.nan)
    name=str(row.get("task_name","")).lower()
    finish=row.get("finish",pd.NaT)
    status=str(row.get("status_code","")).lower()
    completed=("complete" in status and "not" not in status)
    if completed: return "Completed"
    if pd.notna(tf) and tf<0:
        if "lpo" in name or "purchase order" in name or str(row.get("task_code","")).upper().startswith("PO-"):
            return "IMMEDIATE LPO REQUIRED"
        return "CRITICAL / DELAY"
    if pd.notna(finish) and finish < dd:
        return "OVERDUE"
    if "delivery" in name and pd.notna(finish) and finish <= dd+pd.Timedelta(days=14):
        return "DELIVERY DUE"
    if "lpo" in name or "purchase order" in name or str(row.get("task_code","")).upper().startswith("PO-"):
        if pd.notna(finish) and finish <= dd+pd.Timedelta(days=14):
            return "LPO DUE"
    return "Monitor"

def fig_to_png(fig):
    try:
        return fig.to_image(format="png", scale=2)
    except Exception:
        return None

# ------------------------------------------------------------------
# Sidebar inputs
# ------------------------------------------------------------------
st.sidebar.title("Project Controls Hub")
st.sidebar.caption("v0.5 WPR MASTER REPORTING · JRD Project Controls")

project_db_file = st.sidebar.file_uploader("1. Project Database / WPR Inputs", type=["xlsx"], key="project_db")
if DB_TEMPLATE.exists():
    st.sidebar.download_button("Download Project Database Template", DB_TEMPLATE.read_bytes(), "JRD_Project_Data_Input_v1.xlsx", use_container_width=True)
baseline_file = st.sidebar.file_uploader("2. Approved Baseline", type=["xlsx"], key="baseline")
update_file = st.sidebar.file_uploader("3. Current Update", type=["xlsx"], key="update")
resource_file = st.sidebar.file_uploader("4. Resource / Cost Export for Actuals", type=["xlsx"], key="resource")

st.sidebar.markdown("### Planned Curves")
planned_cost_file = st.sidebar.file_uploader("5. Planned Cost / Cash Flow Excel", type=["xlsx"], key="planned_cost")
planned_unit_file = st.sidebar.file_uploader("6. Planned Unit / MHR Excel", type=["xlsx"], key="planned_unit")
planned_manpower_file = st.sidebar.file_uploader("7. Planned Manpower Histogram Excel", type=["xlsx"], key="planned_mph")

today_default = date.today()
data_date_input = st.sidebar.date_input("Data Date", value=today_default, format="DD/MM/YYYY")
data_date = pd.Timestamp(data_date_input)

lookahead_weeks = st.sidebar.selectbox("Lookahead Horizon", [2,4,6,8], index=1)
hours_per_day = st.sidebar.number_input("Working Hours / Day", 1.0, 24.0, 8.0, .5)
days_per_week = st.sidebar.number_input("Working Days / Week", 1.0, 7.0, 6.0, .5)
near_tf = st.sidebar.number_input("Near-Critical Threshold (days)", 1, 90, 20, 1)

st.sidebar.markdown("---")
st.sidebar.caption("Code calculates · AI interprets")

# Progress photos
progress_photos = st.sidebar.file_uploader(
    "8. Progress Photos",
    type=["png","jpg","jpeg"],
    accept_multiple_files=True
)

project_setup, project_db_sheets = read_project_database(project_db_file)
if project_setup:
    db_dd = get_setup(project_setup, "Data Date", None)
    if db_dd is not None and pd.notna(pd.to_datetime(db_dd,errors="coerce")):
        data_date = pd.Timestamp(pd.to_datetime(db_dd))
    near_tf = int(get_setup(project_setup, "Near Critical Threshold (days)", near_tf) or near_tf)
    lookahead_weeks = int(get_setup(project_setup, "Lookahead Weeks", lookahead_weeks) or lookahead_weeks)
    hours_per_day = float(get_setup(project_setup, "Working Hours / Day", hours_per_day) or hours_per_day)
    days_per_week = float(get_setup(project_setup, "Working Days / Week", days_per_week) or days_per_week)

# ------------------------------------------------------------------
# Cover
# ------------------------------------------------------------------
if COVER_IMG.exists():
    cb64=img_b64(COVER_IMG)
    st.markdown(f"""
    <div class="cover-card" style="height:360px;background:url('data:image/png;base64,{cb64}') center/cover;">
      <div class="cover-overlay">
        <div class="cover-title">{str(get_setup(project_setup, "Project Name", "JUMEIRAH RETAIL DEVELOPMENT"))}</div>
        <div class="cover-sub">Owner / Employer: {str(get_setup(project_setup, "Owner / Employer", "ALTA"))} &nbsp;&nbsp; | &nbsp;&nbsp; Contractor: {str(get_setup(project_setup, "Contractor", "DDDC"))} &nbsp;&nbsp; | &nbsp;&nbsp; Consultant: {str(get_setup(project_setup, "Consultant", "DAR"))}</div>
      </div>
    </div>
    """, unsafe_allow_html=True)
else:
    st.title(str(get_setup(project_setup, "Project Name", "JUMEIRAH RETAIL DEVELOPMENT")))
    st.caption(f"Owner / Employer: {get_setup(project_setup, 'Owner / Employer', 'ALTA')} | Contractor: {get_setup(project_setup, 'Contractor', 'DDDC')} | Consultant: {get_setup(project_setup, 'Consultant', 'DAR')}")

if baseline_file is None:
    st.info("Upload the Approved Baseline Primavera Excel export to start.")
    st.stop()

try:
    base_raw, base_pred, base_opt = load_schedule(baseline_file)
    base = enrich(base_raw, base_pred)
except Exception as e:
    st.error(f"Baseline could not be read: {e}")
    st.stop()

curr=None; curr_pred=None; curr_opt={}
if update_file:
    try:
        cr, cp, co = load_schedule(update_file)
        curr=enrich(cr,cp); curr_pred=cp; curr_opt=co
    except Exception as e:
        st.error(f"Current update could not be read: {e}")

# Resource/cost
assignments=None
if resource_file:
    try:
        x=pd.ExcelFile(resource_file)
        for s in x.sheet_names:
            su=s.upper().replace("_"," ").strip()
            if su in {"TASKRSRC","RESOURCEASSIGNMENT","RESOURCE ASSIGNMENT","RESASSIGN","ASSIGNMENTS"}:
                assignments=pd.read_excel(x,sheet_name=s); break
        if assignments is None:
            # best effort: first sheet
            assignments=pd.read_excel(x,sheet_name=x.sheet_names[0])
    except Exception as e:
        st.warning(f"Resource/Cost workbook could not be read: {e}")
if assignments is None:
    assignments = curr_opt.get("assignments") if curr_opt.get("assignments") is not None else base_opt.get("assignments")
assign = assignment_to_activity(assignments)

bf=project_facts(base)
cf=project_facts(curr) if curr is not None else None
forecast_finish = cf["finish"] if cf else bf["finish"]
forecast_move = (pd.Timestamp(forecast_finish).normalize()-pd.Timestamp(bf["finish"]).normalize()).days if pd.notna(forecast_finish) and pd.notna(bf["finish"]) else np.nan
duration_days = (bf["finish"].normalize()-bf["start"].normalize()).days if pd.notna(bf["start"]) and pd.notna(bf["finish"]) else np.nan
elapsed_days = max(0,(data_date.normalize()-bf["start"].normalize()).days) if pd.notna(bf["start"]) else np.nan
remaining_days = max(0,(forecast_finish.normalize()-data_date.normalize()).days) if pd.notna(forecast_finish) else np.nan
elapsed_pct = elapsed_days/duration_days*100 if duration_days and duration_days>0 else np.nan

# Build curves
# Planned values come from user-uploaded planned sheets when provided.
# Actual values come from each update/resource export; missing actuals are shown as zero rather than fabricated.
cost_curve=pd.DataFrame(); unit_curve=pd.DataFrame(); manpower_plan_weekly=pd.DataFrame(); manpower_plan_monthly=pd.DataFrame()

# Planned Cost: user sheet has priority, resource assignment planned cost is fallback.
planned_cost_curve = parse_planned_curve_excel(planned_cost_file, preferred=("PLANNED","COST","W"))
if planned_cost_curve.empty and not assign.empty and "planned_cost" in assign.columns:
    pc=timephase_assign(assign,"planned_cost",data_date,"MS")
    if not pc.empty:
        planned_cost_curve=pc.rename(columns={"planned_cost":"planned"})

# Actual Cost: from resource/cost export; fallback from current update actual cost fields if present.
actual_cost_curve = pd.DataFrame(columns=["date","actual"])
if not assign.empty and "actual_cost" in assign.columns:
    ac=timephase_assign(assign,"actual_cost",data_date,"MS")
    if not ac.empty:
        actual_cost_curve=ac.rename(columns={"actual_cost":"actual"})
if actual_cost_curve.empty and curr is not None:
    actual_cost_curve=actual_curve_from_update(curr,"actual_cost",data_date,"MS")
cost_curve=combine_planned_actual(planned_cost_curve, actual_cost_curve)

# Planned Unit/MHR: user sheet has priority, resource assignment planned units is fallback.
planned_unit_curve = parse_planned_curve_excel(planned_unit_file, preferred=("PLANNED","UNIT","MHR","W"))
if planned_unit_curve.empty and not assign.empty and "planned_units" in assign.columns:
    pu=timephase_assign(assign,"planned_units",data_date,"W")
    if not pu.empty:
        planned_unit_curve=pu.rename(columns={"planned_units":"planned"})

# Actual Unit/MHR: from resource assignments; fallback from current update actual_units if present.
actual_unit_curve = pd.DataFrame(columns=["date","actual"])
if not assign.empty and "actual_units" in assign.columns:
    au=timephase_assign(assign,"actual_units",data_date,"W")
    if not au.empty:
        actual_unit_curve=au.rename(columns={"actual_units":"actual"})
if actual_unit_curve.empty and curr is not None:
    actual_unit_curve=actual_curve_from_update(curr,"actual_units",data_date,"W")
unit_curve=combine_planned_actual(planned_unit_curve, actual_unit_curve)

# Planned manpower from separate manpower histogram; actual manpower entered manually in app.
manpower_plan_weekly = parse_manpower_plan_excel(planned_manpower_file, monthly=False)
manpower_plan_monthly = parse_manpower_plan_excel(planned_manpower_file, monthly=True)

# Executive progress KPIs
def progress_from_curve(curve):
    if curve.empty: return (np.nan,np.nan,np.nan)
    c=curve[curve["date"]<=data_date]
    if c.empty: return (0.0,0.0,0.0)
    p=float(c["cum_planned_pct"].iloc[-1]) if "cum_planned_pct" in c else np.nan
    a=float(c["cum_actual_pct"].iloc[-1]) if "cum_actual_pct" in c else np.nan
    return p,a,a-p if pd.notna(p) and pd.notna(a) else np.nan
cost_plan,cost_actual,cost_var=progress_from_curve(cost_curve)
mhr_plan,mhr_actual,mhr_var=progress_from_curve(unit_curve)

tabs = st.tabs([
    "Executive Dashboard","Engineering","Procurement","Progress Curves",
    "Cash Flow","Manpower","Milestones","Schedule Health","4-Week Lookahead",
    "Progress Photos","T&C / Handover","Report Center"
])

# ------------------------------------------------------------------
# Executive Dashboard
# ------------------------------------------------------------------
with tabs[0]:
    st.markdown('<div class="section-title">Executive Project Summary</div>', unsafe_allow_html=True)
    a,b = st.columns([1,2])
    with a:
        summary = pd.DataFrame({
            "Description":["Contract Commencement","Contract Completion","Forecast Completion","Total Duration (days)",
                           "Time Elapsed (days)","Remaining Time (days)","Time Elapsed (%)","Data Date","Schedule Status"],
            "Status":[fmt_date(pd.to_datetime(get_setup(project_setup,"Contract Commencement",bf["start"]),errors="coerce")),fmt_date(pd.to_datetime(get_setup(project_setup,"Contract Completion",bf["finish"]),errors="coerce")),fmt_date(forecast_finish),
                      int(duration_days) if pd.notna(duration_days) else None,
                      int(elapsed_days) if pd.notna(elapsed_days) else None,
                      int(remaining_days) if pd.notna(remaining_days) else None,
                      f"{elapsed_pct:.1f}%" if pd.notna(elapsed_pct) else "—",
                      fmt_date(data_date),
                      "Behind Programme" if pd.notna(forecast_move) and forecast_move>0 else "On / Ahead"]
        })
        st.dataframe(summary, use_container_width=True, hide_index=True)
    with b:
        c=st.columns(6)
        c[0].metric("Forecast Movement", f"{forecast_move:+.0f} d" if pd.notna(forecast_move) else "—")
        c[1].metric("Critical Activities", f"{bf['critical']:,}")
        c[2].metric("Negative Float", f"{int(((curr if curr is not None else base).total_float_days<0).sum()):,}")
        c[3].metric("MHR Planned", f"{mhr_plan:.2f}%" if pd.notna(mhr_plan) else "N/A")
        c[4].metric("MHR Actual", f"{mhr_actual:.2f}%" if pd.notna(mhr_actual) else "N/A")
        c[5].metric("Cost Actual", f"{cost_actual:.2f}%" if pd.notna(cost_actual) else "N/A")

        # Progress summary tables
        ptab = pd.DataFrame([
            ["MHR",mhr_plan,mhr_actual,mhr_var],
            ["Cost",cost_plan,cost_actual,cost_var],
        ], columns=["Weightage Basis","Cumulative Plan %","Cumulative Actual %","Variance %"])
        st.dataframe(
            ptab.style.format({"Cumulative Plan %":"{:.2f}%","Cumulative Actual %":"{:.2f}%","Variance %":"{:.2f}%"}).map(
                lambda v: f"color:{BRAND_RED};font-weight:700" if isinstance(v,(float,np.floating)) and v<0 else "",
                subset=["Variance %"]
            ),
            use_container_width=True, hide_index=True
        )

    c1,c2=st.columns(2)
    with c1:
        # Duration chart
        dd=pd.DataFrame({"Metric":["Original Duration","At Completion Duration","Time Elapsed"],
                         "Days":[duration_days, duration_days+(forecast_move if pd.notna(forecast_move) else 0), elapsed_days]})
        fig=px.bar(dd,x="Days",y="Metric",orientation="h",text="Days",title="Project Duration")
        fig.update_layout(height=320,xaxis_title="Days",yaxis_title="")
        st.plotly_chart(fig,use_container_width=True)
    with c2:
        # Management alerts
        active = curr if curr is not None else base
        engrisk=active[active["category"].isin(["Prequalification","Shop Drawings","Material Submittals"]) & (active["total_float_days"]<=near_tf)]
        procrisk=active[(active["category"]=="Procurement") & (active["total_float_days"]<=near_tf)]
        alerts=[
            f"{int((active.total_float_days<0).sum())} activities with negative float",
            f"{len(engrisk)} engineering items at/near critical",
            f"{len(procrisk)} procurement items at/near critical",
            f"Forecast completion movement: {forecast_move:+.0f} days" if pd.notna(forecast_move) else "Forecast movement unavailable",
            f"{lookahead_weeks}-week lookahead starts from {fmt_date(data_date)}",
        ]
        st.markdown("### Management Alerts")
        for x in alerts:
            st.write("•",x)

# ------------------------------------------------------------------
# Engineering
# ------------------------------------------------------------------
with tabs[1]:
    st.markdown('<div class="section-title">Engineering Control · Major Package Wise</div>', unsafe_allow_html=True)
    active=curr if curr is not None else base
    eng=active[active["category"].isin(["Prequalification","Shop Drawings","Material Submittals"])].copy()
    k=st.columns(5)
    k[0].metric("Engineering Items",len(eng))
    k[1].metric("PQ",int((eng.category=="Prequalification").sum()))
    k[2].metric("Shop Drawings",int((eng.category=="Shop Drawings").sum()))
    k[3].metric("Material Submittals",int((eng.category=="Material Submittals").sum()))
    k[4].metric("Negative Float",int((eng.total_float_days<0).sum()))

    matrix = eng.pivot_table(
        index="package", columns="category", values="finish", aggfunc="max"
    ).reset_index()
    st.subheader("Package Matrix")
    st.dataframe(matrix,use_container_width=True,hide_index=True)

    eng["Risk / Status"]=eng["total_float_days"].apply(status_from_float)
    eng["Days to Due"]=(eng["finish"]-data_date).dt.days
    st.subheader("Engineering Action Register")
    show=eng[["package","task_code","task_name","category","finish","total_float_days","Days to Due","Risk / Status"]].sort_values(
        ["total_float_days","finish"],na_position="last"
    )
    styled=show.style.format({"total_float_days":"{:.0f} d","Days to Due":"{:.0f} d"}).map(style_float,subset=["total_float_days"])
    st.dataframe(styled,use_container_width=True,hide_index=True)

# ------------------------------------------------------------------
# Procurement
# ------------------------------------------------------------------
with tabs[2]:
    st.markdown('<div class="section-title">Procurement Action Dashboard</div>', unsafe_allow_html=True)
    active=curr if curr is not None else base
    proc=active[active["category"]=="Procurement"].copy()
    proc["Action"]=proc.apply(procurement_action,axis=1,dd=data_date)
    proc["Days to Finish"]=(proc["finish"]-data_date).dt.days
    priority_order={"IMMEDIATE LPO REQUIRED":0,"CRITICAL / DELAY":1,"OVERDUE":2,"LPO DUE":3,"DELIVERY DUE":4,"Monitor":5,"Completed":6}
    proc["Priority"]=proc["Action"].map(priority_order).fillna(9)
    k=st.columns(6)
    labels=["IMMEDIATE LPO REQUIRED","CRITICAL / DELAY","OVERDUE","LPO DUE","DELIVERY DUE","Completed"]
    for i,lab in enumerate(labels):
        k[i].metric(lab.title(),int((proc["Action"]==lab).sum()))
    st.subheader("Priority Action Register")
    pshow=proc.sort_values(["Priority","total_float_days","finish"])[
        ["task_code","task_name","finish","total_float_days","Days to Finish","Action"]
    ]
    st.dataframe(
        pshow.style.format({"total_float_days":"{:.0f} d","Days to Finish":"{:.0f} d"}).map(style_float,subset=["total_float_days"]),
        use_container_width=True, hide_index=True
    )

# ------------------------------------------------------------------
# Progress Curves
# ------------------------------------------------------------------
with tabs[3]:
    st.markdown('<div class="section-title">Progress Curves</div>', unsafe_allow_html=True)
    c1,c2=st.columns(2)
    with c1:
        st.subheader("MHR / Unit S-Curve")
        if unit_curve.empty:
            st.warning("Resource Assignment units are required.")
        else:
            fig=go.Figure()
            fig.add_trace(go.Scatter(x=unit_curve.date,y=unit_curve.cum_planned_pct,name="Planned",mode="lines"))
            fig.add_trace(go.Scatter(x=unit_curve.date,y=unit_curve.cum_actual_pct,name="Actual",mode="lines"))
            fig.update_layout(height=420,yaxis_title="Cumulative %",xaxis_title="")
            st.plotly_chart(fig,use_container_width=True)
    with c2:
        st.subheader("Cost S-Curve")
        if cost_curve.empty:
            st.warning("Cost loading is required.")
        else:
            fig=go.Figure()
            fig.add_trace(go.Scatter(x=cost_curve.date,y=cost_curve.cum_planned_pct,name="Planned",mode="lines"))
            fig.add_trace(go.Scatter(x=cost_curve.date,y=cost_curve.cum_actual_pct,name="Actual",mode="lines"))
            fig.update_layout(height=420,yaxis_title="Cumulative %",xaxis_title="")
            st.plotly_chart(fig,use_container_width=True)

# ------------------------------------------------------------------
# Cash Flow
# ------------------------------------------------------------------
with tabs[4]:
    st.markdown('<div class="section-title">Cash Flow · Planned vs Actual</div>', unsafe_allow_html=True)
    if cost_curve.empty:
        st.warning("Planned and Actual Cost fields are required from Resource / Cost Assignments.")
    else:
        fig=go.Figure()
        fig.add_bar(x=cost_curve.date,y=cost_curve.planned,name="Monthly Planned")
        fig.add_bar(x=cost_curve.date,y=cost_curve.actual,name="Monthly Actual")
        fig.add_trace(go.Scatter(x=cost_curve.date,y=cost_curve.cum_planned,name="Cumulative Planned",yaxis="y2",mode="lines+markers"))
        fig.add_trace(go.Scatter(x=cost_curve.date,y=cost_curve.cum_actual,name="Cumulative Actual",yaxis="y2",mode="lines+markers"))
        fig.update_layout(barmode="group",height=500,yaxis_title="Monthly Cost",
                          yaxis2=dict(title="Cumulative Cost",overlaying="y",side="right"),xaxis_title="")
        st.plotly_chart(fig,use_container_width=True)
        st.dataframe(cost_curve,use_container_width=True,hide_index=True)

# ------------------------------------------------------------------
# Manpower
# ------------------------------------------------------------------
with tabs[5]:
    st.markdown('<div class="section-title">Manpower Histogram · Planned from Excel / Actual Manual</div>', unsafe_allow_html=True)
    st.caption("Planned manpower comes from the manpower histogram Excel. Actual manpower is entered from the project timesheet.")
    if manpower_plan_weekly.empty:
        st.warning("Upload the Planned Manpower Histogram Excel to activate planned manpower charts.")
        mp = pd.DataFrame({"date": pd.date_range(data_date, periods=8, freq="W"), "planned": 0.0})
    else:
        mp = manpower_plan_weekly.copy()

    actual_seed = mp[["date","planned"]].copy()
    actual_seed["date"] = pd.to_datetime(actual_seed["date"])
    actual_seed["actual"] = 0.0
    db_mp = db_sheet(project_db_sheets, "Actual Manpower")
    if not db_mp.empty and "Date" in db_mp.columns:
        db_mp["Date"] = pd.to_datetime(db_mp["Date"], errors="coerce")
        if "Total" not in db_mp.columns:
            parts=[c for c in ["Staff","Civil","MEP","Finishes","Subcontractors"] if c in db_mp.columns]
            if parts:
                db_mp["Total"] = db_mp[parts].apply(pd.to_numeric,errors="coerce").fillna(0).sum(axis=1)
        db_mp["Total"] = pd.to_numeric(db_mp.get("Total"), errors="coerce")
        db_week = db_mp.dropna(subset=["Date"]).set_index("Date")["Total"].resample("W").mean().reset_index().rename(columns={"Date":"date","Total":"actual_db"})
        actual_seed = pd.merge(actual_seed, db_week, on="date", how="left")
        actual_seed["actual"] = actual_seed["actual_db"].fillna(0.0)
        actual_seed = actual_seed.drop(columns=["actual_db"])
    actual_seed["date"] = actual_seed["date"].dt.date
    st.markdown("### Actual Manpower Input from Timesheet")
    edited = st.data_editor(
        actual_seed,
        column_config={
            "date": st.column_config.DateColumn("Week / Month"),
            "planned": st.column_config.NumberColumn("Planned Manpower", disabled=True, format="%.0f"),
            "actual": st.column_config.NumberColumn("Actual Manpower", min_value=0.0, step=1.0, format="%.0f"),
        },
        hide_index=True,
        use_container_width=True,
        key="actual_manpower_editor"
    )
    mp_chart = edited.copy()
    mp_chart["date"] = pd.to_datetime(mp_chart["date"])
    mp_chart["variance"] = mp_chart["actual"] - mp_chart["planned"]

    c1,c2=st.columns(2)
    with c1:
        fig=px.bar(mp_chart,x="date",y=["planned","actual"],barmode="group",title="Weekly/Monthly Planned vs Actual Manpower")
        fig.update_layout(height=430,xaxis_title="",yaxis_title="Headcount")
        st.plotly_chart(fig,use_container_width=True)
    with c2:
        fig=px.bar(mp_chart,x="date",y="variance",title="Manpower Variance")
        fig.update_layout(height=430,xaxis_title="",yaxis_title="Actual - Planned")
        st.plotly_chart(fig,use_container_width=True)

    if not unit_curve.empty:
        st.markdown("### MHR / Labor Units")
        unit_curve["planned_manpower_equivalent"]=unit_curve["planned"]/(hours_per_day*days_per_week)
        unit_curve["actual_manpower_equivalent"]=unit_curve["actual"]/(hours_per_day*days_per_week)
        fig=px.bar(unit_curve,x="date",y=["planned","actual"],barmode="group",title="Weekly Planned vs Actual Manhours / Labor Units")
        fig.update_layout(height=380,xaxis_title="",yaxis_title="Manhours / Labor Units")
        st.plotly_chart(fig,use_container_width=True)
        st.caption(f"Equivalent headcount conversion: weekly labor hours ÷ ({hours_per_day:g} h/day × {days_per_week:g} days/week).")

# ------------------------------------------------------------------
# Milestones
# ------------------------------------------------------------------
with tabs[6]:
    st.markdown('<div class="section-title">Key Milestones</div>', unsafe_allow_html=True)
    active=curr if curr is not None else base
    ms=active[active["duration_days"].fillna(-1)==0].copy()
    ms["Status"]=ms["total_float_days"].apply(status_from_float)
    st.dataframe(
        ms[["task_code","task_name","finish","total_float_days","Status"]].sort_values("finish")
          .style.format({"total_float_days":"{:.0f} d"}).map(style_float,subset=["total_float_days"]),
        use_container_width=True,hide_index=True
    )

# ------------------------------------------------------------------
# Schedule health
# ------------------------------------------------------------------
with tabs[7]:
    st.markdown('<div class="section-title">Schedule Health</div>', unsafe_allow_html=True)
    active=curr if curr is not None else base
    pred=curr_pred if curr_pred is not None else base_pred
    pc=pred.groupby("task_id").size(); sc=pred.groupby("pred_task_id").size()
    nonms=active[active["duration_days"].fillna(0)>0]
    op=nonms[~nonms.task_code.isin(pc.index)]
    os=nonms[~nonms.task_code.isin(sc.index)]
    c=st.columns(5)
    c[0].metric("No Predecessor",len(op))
    c[1].metric("No Successor",len(os))
    c[2].metric("Critical / Zero Float",int((active.total_float_days<=0).sum()))
    c[3].metric("Negative Float",int((active.total_float_days<0).sum()))
    c[4].metric(f"Near Critical ≤ {near_tf}d",int(((active.total_float_days>0)&(active.total_float_days<=near_tf)).sum()))
    low=active.sort_values("total_float_days").head(100)
    st.dataframe(
        low[["task_code","task_name","finish","total_float_days","category"]]
           .style.format({"total_float_days":"{:.0f} d"}).map(style_float,subset=["total_float_days"]),
        use_container_width=True,hide_index=True
    )

# ------------------------------------------------------------------
# Lookahead
# ------------------------------------------------------------------
with tabs[8]:
    st.markdown(f'<div class="section-title">{lookahead_weeks}-Week Lookahead from Data Date</div>', unsafe_allow_html=True)
    active=curr if curr is not None else base
    la=date_window(active,data_date,lookahead_weeks)
    st.write(f"**Data Date:** {fmt_date(data_date)}  |  **Window End:** {fmt_date(data_date+pd.Timedelta(weeks=lookahead_weeks))}")
    st.dataframe(
        la[["task_code","task_name","category","start","finish","total_float_days","status_code"]]
          .style.format({"total_float_days":"{:.0f} d"}).map(style_float,subset=["total_float_days"]),
        use_container_width=True,hide_index=True
    )

# ------------------------------------------------------------------
# Photos
# ------------------------------------------------------------------
with tabs[9]:
    st.markdown('<div class="section-title">Progress Photos</div>', unsafe_allow_html=True)
    if not progress_photos:
        st.info("Upload weekly site progress photos from the sidebar.")
    else:
        cols=st.columns(3)
        for i,p in enumerate(progress_photos):
            with cols[i%3]:
                st.image(p,use_container_width=True)
                st.caption(p.name)

# ------------------------------------------------------------------
# T&C
# ------------------------------------------------------------------
with tabs[10]:
    st.markdown('<div class="section-title">T&C / Handover</div>', unsafe_allow_html=True)
    active=curr if curr is not None else base
    tch=active[active.category.isin(["T&C","Handover"])].sort_values("start")
    st.dataframe(
        tch[["task_code","task_name","category","start","finish","total_float_days"]]
          .style.format({"total_float_days":"{:.0f} d"}).map(style_float,subset=["total_float_days"]),
        use_container_width=True,hide_index=True
    )

# ------------------------------------------------------------------
# Report generators
# ------------------------------------------------------------------
def _curve_resample(curve, freq="MS"):
    if curve is None or curve.empty:
        return pd.DataFrame(columns=["date","planned","actual","cum_planned","cum_actual","cum_planned_pct","cum_actual_pct"])
    d=curve.copy()
    d["date"]=pd.to_datetime(d["date"],errors="coerce")
    d=d.dropna(subset=["date"]).set_index("date")
    for col in ["planned","actual"]:
        if col not in d.columns: d[col]=0.0
    out=d[["planned","actual"]].resample(freq).sum().reset_index()
    out["cum_planned"]=out["planned"].cumsum(); out["cum_actual"]=out["actual"].cumsum()
    pt=out["planned"].sum()
    out["cum_planned_pct"]=np.where(pt!=0,out["cum_planned"]/pt*100,0.0)
    out["cum_actual_pct"]=np.where(pt!=0,out["cum_actual"]/pt*100,0.0)
    return out

def _history_snapshot():
    active=curr if curr is not None else base
    return build_history_curves(
        project_db_sheets, data_date,
        current_mhr_actual=mhr_actual, current_cost_actual=cost_actual,
        current_mhr_plan=mhr_plan, current_cost_plan=cost_plan
    )

def _current_week_summary(history):
    if history is None or history.empty:
        return {k:np.nan for k in ["mhr_wp","mhr_wa","mhr_cp","mhr_ca","mhr_var","cost_wp","cost_wa","cost_cp","cost_ca","cost_var"]}
    r=history.sort_values("Data Date").iloc[-1]
    return {
        "mhr_wp":r.get("MHR Weekly Plan %",np.nan), "mhr_wa":r.get("MHR Weekly Actual %",np.nan),
        "mhr_cp":r.get("MHR Cumulative Plan %",np.nan), "mhr_ca":r.get("MHR Cumulative Actual %",np.nan),
        "mhr_var":r.get("MHR Cumulative Actual %",np.nan)-r.get("MHR Cumulative Plan %",np.nan),
        "cost_wp":r.get("Cost Weekly Plan %",np.nan), "cost_wa":r.get("Cost Weekly Actual %",np.nan),
        "cost_cp":r.get("Cost Cumulative Plan %",np.nan), "cost_ca":r.get("Cost Cumulative Actual %",np.nan),
        "cost_var":r.get("Cost Cumulative Actual %",np.nan)-r.get("Cost Cumulative Plan %",np.nan),
    }

def _db_text(section, default="Data not provided"):
    df=db_sheet(project_db_sheets,"Weekly Narrative")
    if df.empty or "Section" not in df.columns:
        return default
    hit=df[df["Section"].astype(str).str.strip().str.lower()==section.strip().lower()]
    if hit.empty: return default
    col="Narrative / Details" if "Narrative / Details" in hit.columns else hit.columns[-1]
    v=hit.iloc[0][col]
    return default if pd.isna(v) or str(v).strip()=="" else str(v)


def _mpl_bytes(fig):
    out=BytesIO(); fig.savefig(out,format="png",dpi=170,bbox_inches="tight",facecolor="white"); plt.close(fig); out.seek(0); return out.getvalue()

def _curve_png(curve,title):
    if curve is None or curve.empty:return None
    fig,ax=plt.subplots(figsize=(12.5,5.5))
    ax.plot(curve["date"],curve["cum_planned_pct"],marker="o",linewidth=2.4,label="Planned",color=BRAND_GREEN)
    ax.plot(curve["date"],curve["cum_actual_pct"],marker="o",linewidth=2.4,label="Actual",color=BRAND_YELLOW)
    ax.set_title(title,fontsize=14,fontweight="bold");ax.set_ylabel("Cumulative Progress %");ax.grid(True,alpha=.22);ax.legend(loc="upper left",ncol=2)
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b-%y"));fig.autofmt_xdate(rotation=35)
    fig.tight_layout();return _mpl_bytes(fig)

def _cashflow_png(curve,title):
    if curve is None or curve.empty:return None
    fig,ax=plt.subplots(figsize=(12.5,5.5));x=pd.to_datetime(curve["date"])
    width=12 if len(curve)<24 else 3
    ax.bar(x-pd.to_timedelta(width/2,unit="D"),curve["planned"],width=width,label="Planned Period",color=BRAND_GREEN)
    ax.bar(x+pd.to_timedelta(width/2,unit="D"),curve["actual"],width=width,label="Actual Period",color=BRAND_YELLOW)
    ax2=ax.twinx();ax2.plot(x,curve["cum_planned"],linewidth=2.2,label="Cum Planned",color="#246B8E");ax2.plot(x,curve["cum_actual"],linewidth=2.2,label="Cum Actual",color=BRAND_RED)
    ax.set_title(title,fontsize=14,fontweight="bold");ax.set_ylabel("Period Value");ax2.set_ylabel("Cumulative Value");ax.grid(True,axis="y",alpha=.2)
    lines,labels=ax.get_legend_handles_labels(); lines2,labels2=ax2.get_legend_handles_labels(); ax.legend(lines+lines2,labels+labels2,loc="upper left",ncol=4,fontsize=8)
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b-%y"));fig.autofmt_xdate(rotation=35);fig.tight_layout();return _mpl_bytes(fig)

def _manpower_png(mp,title):
    if mp is None or mp.empty:return None
    fig,ax=plt.subplots(figsize=(12.5,5.5));x=pd.to_datetime(mp["date"]);idx=np.arange(len(x));w=.38
    ax.bar(idx-w/2,pd.to_numeric(mp["planned"],errors="coerce").fillna(0),w,label="Planned",color=BRAND_GREEN)
    ax.bar(idx+w/2,pd.to_numeric(mp["actual"],errors="coerce").fillna(0),w,label="Actual",color=BRAND_YELLOW)
    step=max(1,len(idx)//14);ax.set_xticks(idx[::step]);ax.set_xticklabels([d.strftime("%d-%b-%y") for d in x.iloc[::step]],rotation=40,ha="right")
    ax.set_title(title,fontsize=14,fontweight="bold");ax.set_ylabel("Headcount");ax.grid(True,axis="y",alpha=.2);ax.legend();fig.tight_layout();return _mpl_bytes(fig)

def _compare_png(labels,planned,actual,title):
    fig,ax=plt.subplots(figsize=(7.5,4.8));idx=np.arange(len(labels));w=.36
    ax.bar(idx-w/2,planned,w,label="Planned",color=BRAND_GREEN);ax.bar(idx+w/2,actual,w,label="Actual",color=BRAND_YELLOW)
    ax.set_xticks(idx);ax.set_xticklabels(labels);ax.set_title(title,fontweight="bold");ax.set_ylabel("%");ax.legend();ax.grid(True,axis="y",alpha=.2);fig.tight_layout();return _mpl_bytes(fig)

def _fig_png(fig, width=1500, height=780):
    if fig is None: return None
    try:
        fig.update_layout(template="plotly_white",font=dict(family="Arial",size=14),
                          legend=dict(orientation="h",yanchor="bottom",y=1.02,xanchor="right",x=1),
                          margin=dict(l=65,r=45,t=55,b=55))
        return fig.to_image(format="png",width=width,height=height,scale=2)
    except Exception:
        return None

def _curve_fig(curve, title, planned_label="Planned", actual_label="Actual", weekly=False):
    if curve is None or curve.empty: return None
    f=go.Figure()
    f.add_trace(go.Scatter(x=curve["date"],y=curve["cum_planned_pct"],name=planned_label,mode="lines+markers",line=dict(width=3,color=BRAND_GREEN)))
    f.add_trace(go.Scatter(x=curve["date"],y=curve["cum_actual_pct"],name=actual_label,mode="lines+markers",line=dict(width=3,color=BRAND_YELLOW)))
    f.update_layout(title=title,yaxis_title="Cumulative Progress %",xaxis_title="",yaxis=dict(range=[0,max(100,float(curve[["cum_planned_pct","cum_actual_pct"]].max().max())*1.08)]))
    return f

def _cashflow_fig(curve, title):
    if curve is None or curve.empty:return None
    f=go.Figure()
    f.add_bar(x=curve["date"],y=curve["planned"],name="Planned Period",marker_color=BRAND_GREEN)
    f.add_bar(x=curve["date"],y=curve["actual"],name="Actual Period",marker_color=BRAND_YELLOW)
    f.add_trace(go.Scatter(x=curve["date"],y=curve["cum_planned"],name="Cumulative Planned",mode="lines+markers",yaxis="y2",line=dict(color="#246B8E",width=3)))
    f.add_trace(go.Scatter(x=curve["date"],y=curve["cum_actual"],name="Cumulative Actual",mode="lines+markers",yaxis="y2",line=dict(color=BRAND_RED,width=3)))
    f.update_layout(title=title,barmode="group",yaxis_title="Period Value",yaxis2=dict(title="Cumulative",overlaying="y",side="right"),xaxis_title="")
    return f

def _manpower_fig(mp, title):
    if mp is None or mp.empty:return None
    f=go.Figure()
    f.add_bar(x=mp["date"],y=mp["planned"],name="Planned",marker_color=BRAND_GREEN)
    f.add_bar(x=mp["date"],y=mp["actual"],name="Actual",marker_color=BRAND_YELLOW)
    f.update_layout(title=title,barmode="group",yaxis_title="Headcount",xaxis_title="")
    return f


def make_updated_project_database():
    """Append/replace the current update snapshot while preserving the user's workbook formatting."""
    try:
        raw = project_db_file.getvalue() if project_db_file is not None else DB_TEMPLATE.read_bytes()
        wb = load_workbook(BytesIO(raw))
        if "Update History" not in wb.sheetnames:
            ws = wb.create_sheet("Update History")
            ws.append(["Data Date","Report No.","MHR Cumulative Plan %","MHR Cumulative Actual %","Cost Cumulative Plan %","Cost Cumulative Actual %","Forecast Completion","Negative Float Activities","Notes"])
        ws=wb["Update History"]
        target=None
        for r in range(2,ws.max_row+1):
            v=ws.cell(r,1).value
            if v and pd.notna(pd.to_datetime(v,errors="coerce")) and pd.Timestamp(pd.to_datetime(v)).normalize()==pd.Timestamp(data_date).normalize():
                target=r;break
        if target is None: target=max(2,ws.max_row+1)
        active=curr if curr is not None else base
        values=[pd.Timestamp(data_date).to_pydatetime(),get_setup(project_setup,"Report Number",None),mhr_plan,mhr_actual,cost_plan,cost_actual,
                pd.Timestamp(forecast_finish).to_pydatetime() if pd.notna(forecast_finish) else None,int((active.total_float_days<0).sum()),"Updated by Project Controls Hub"]
        for col,val in enumerate(values,1): ws.cell(target,col).value=val
        out=BytesIO();wb.save(out);out.seek(0);return out.getvalue()
    except Exception:
        return None

def make_pdf():
    """Client/consultant-grade A3 WPR following the reference report structure."""
    buf=BytesIO(); c=canvas.Canvas(buf,pagesize=A3)
    project=str(get_setup(project_setup,"Project Name","JUMEIRAH RETAIL DEVELOPMENT"))
    owner=str(get_setup(project_setup,"Owner / Employer","ALTA")); consultant=str(get_setup(project_setup,"Consultant","DAR")); contractor=str(get_setup(project_setup,"Contractor","DDDC"))
    report_no=get_setup(project_setup,"Report Number","")
    period_from=pd.to_datetime(get_setup(project_setup,"Reporting Period From",pd.NaT),errors="coerce")
    period_to=pd.to_datetime(get_setup(project_setup,"Reporting Period To",data_date),errors="coerce")
    active=curr if curr is not None else base
    neg_count=int((active.total_float_days<0).sum())
    critical_count=int((active.total_float_days<=0).sum())
    history=_history_snapshot(); wk=_current_week_summary(history)
    page_no=0

    def logos(canvas_obj, W, H, top=18*mm):
        items=[(ALTA_LOGO,18*mm),(DDDC_LOGO,18*mm),(DAR_LOGO,12*mm)]
        x=16*mm
        for path,h in items:
            if path.exists():
                try:
                    im=ImageReader(str(path)); iw,ih=im.getSize(); w=h*iw/ih
                    canvas_obj.drawImage(im,x,H-top-h,width=w,height=h,mask='auto',preserveAspectRatio=True)
                    x+=w+8*mm
                except Exception: pass

    def footer(canvas_obj,W,H):
        nonlocal page_no
        canvas_obj.setStrokeColor(colors.HexColor(BRAND_SILVER)); canvas_obj.line(15*mm,12*mm,W-15*mm,12*mm)
        canvas_obj.setFont("Helvetica",7.5); canvas_obj.setFillColor(colors.HexColor("#5A5A5A"))
        canvas_obj.drawString(15*mm,7*mm,f"{project} | Weekly Progress Report {report_no} | Data Date {fmt_date(data_date)}")
        canvas_obj.drawRightString(W-15*mm,7*mm,f"Page {page_no}")

    def page(title, landscape_mode=False, subtitle=None):
        nonlocal page_no
        if page_no>0: c.showPage()
        size=landscape(A3) if landscape_mode else A3; c.setPageSize(size); W,H=size; page_no+=1
        c.setFillColor(colors.white); c.rect(0,0,W,H,fill=1,stroke=0)
        logos(c,W,H)
        c.setFillColor(colors.HexColor(BRAND_GREEN)); c.rect(0,H-37*mm,W,7*mm,fill=1,stroke=0)
        c.setFillColor(colors.HexColor(BRAND_BLACK)); c.setFont("Helvetica-Bold",18); c.drawString(16*mm,H-48*mm,title)
        if subtitle:
            c.setFont("Helvetica",8.5); c.setFillColor(colors.HexColor("#666666")); c.drawString(16*mm,H-55*mm,subtitle)
        footer(c,W,H)
        return W,H,H-64*mm

    def draw_table(df,x,y,w,h=None,font=7.5,headers=True,negative_cols=None,max_rows=None,col_widths=None):
        if df is None or df.empty:
            c.setFillColor(colors.HexColor("#666666")); c.setFont("Helvetica-Oblique",10); c.drawString(x,y,"Data not provided")
            return y-8*mm
        d=df.copy()
        if max_rows is not None: d=d.head(max_rows)
        vals=[list(map(lambda z:"" if pd.isna(z) else str(z),d.columns))] if headers else []
        for _,r in d.iterrows(): vals.append(["" if pd.isna(v) else str(v) for v in r.tolist()])
        n=max(1,len(d.columns))
        widths=col_widths if col_widths else [w/n]*n
        t=Table(vals,colWidths=widths,repeatRows=1 if headers else 0)
        style=[("BACKGROUND",(0,0),(-1,0),colors.HexColor(BRAND_GREEN)),("TEXTCOLOR",(0,0),(-1,0),colors.white),
               ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),font),
               ("GRID",(0,0),(-1,-1),0.3,colors.HexColor(BRAND_SILVER)),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
               ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F7F7F5")])]
        if negative_cols:
            for idx in negative_cols:
                for rr in range(1,len(vals)):
                    try:
                        txt=str(vals[rr][idx]).replace(' d','').replace('%','').replace(',','')
                        if float(txt)<0: style.append(("TEXTCOLOR",(idx,rr),(idx,rr),colors.HexColor(BRAND_RED))); style.append(("FONTNAME",(idx,rr),(idx,rr),"Helvetica-Bold"))
                    except Exception: pass
        t.setStyle(TableStyle(style)); tw,th=t.wrap(w,h or 9999); t.drawOn(c,x,y-th); return y-th-5*mm

    def draw_chart(fig,x,y,w,h):
        png=fig if isinstance(fig,(bytes,bytearray)) else _fig_png(fig)
        if png:
            try: c.drawImage(ImageReader(BytesIO(png)),x,y-h,width=w,height=h,mask='auto',preserveAspectRatio=True); return
            except Exception: pass
        c.setStrokeColor(colors.HexColor(BRAND_SILVER)); c.rect(x,y-h,w,h,stroke=1,fill=0)
        c.setFont("Helvetica-Oblique",9); c.setFillColor(colors.HexColor("#777777")); c.drawCentredString(x+w/2,y-h/2,"Chart unavailable / data not provided")

    # 1 COVER
    W,H=A3; c.setPageSize(A3); page_no=1
    if COVER_IMG.exists():
        try:c.drawImage(ImageReader(str(COVER_IMG)),0,0,width=W,height=H,mask='auto',preserveAspectRatio=True,anchor='c')
        except Exception: pass
    c.setFillColor(colors.Color(0,0,0,alpha=.55)); c.rect(0,0,W,H,fill=1,stroke=0)
    logos(c,W,H,top=20*mm)
    c.setFillColor(colors.white); c.setFont("Helvetica-Bold",30); c.drawString(22*mm,94*mm,project)
    c.setFont("Helvetica",13); c.drawString(22*mm,82*mm,"WEEKLY PROGRESS REPORT")
    c.setFont("Helvetica-Bold",11); c.drawString(22*mm,69*mm,f"Report No.: {report_no}")
    c.setFont("Helvetica",10); c.drawString(22*mm,59*mm,f"Reporting Period: {fmt_date(period_from)} to {fmt_date(period_to)}")
    c.drawString(22*mm,50*mm,f"Data Date: {fmt_date(data_date)}")
    c.drawString(22*mm,36*mm,f"Owner / Employer: {owner}    |    Consultant: {consultant}    |    Contractor: {contractor}")

    # 2 TOC
    W,H,y=page("Table of Contents",False)
    toc=[
        ["01","Project Overview"],["02","Executive Dashboard"],["03","Progress Summary"],["04","Major Trades Tracker"],
        ["05","Progress Analysis / Narrative"],["06","Labor Force"],["07","AoC & Risk"],["08","Updated S-Curves"],
        ["09","Engineering Control"],["10","Procurement Control"],["11",f"{lookahead_weeks}-Week Lookahead"],
        ["12","Structure & Finishes Trackers"],["13","QC Report"],["14","RFI"],["15","Commercial Status"],["16","Progress Photos"]]
    draw_table(pd.DataFrame(toc,columns=["Section","Description"]),28*mm,y-8*mm,W-56*mm,font=10,col_widths=[25*mm,W-81*mm])

    # 3 Project Overview
    W,H,y=page("1.1 - Project Overview",False)
    setup_rows=[]
    preferred=["Project Name","Owner / Employer","Consultant","Contractor","Contract Commencement","Contract Completion","Contract Value (AED)","Project Manager","Construction Manager","Planning Engineer","Commercial / QS","HSE Manager","QA/QC Manager"]
    for k in preferred:
        v=get_setup(project_setup,k,"Data not provided")
        if "Date" in k or "Commencement" in k or "Completion" in k: v=fmt_date(pd.to_datetime(v,errors="coerce")) if pd.notna(pd.to_datetime(v,errors="coerce")) else str(v)
        setup_rows.append([k,v])
    draw_table(pd.DataFrame(setup_rows,columns=["Description","Status"]),25*mm,y-5*mm,W-50*mm,font=9,col_widths=[78*mm,W-128*mm])

    # 4 Dashboard landscape
    W,H,y=page("Executive Dashboard",True,"Overall project performance, schedule, MHR and cost progress")
    contract_start=pd.to_datetime(get_setup(project_setup,"Contract Commencement",bf["start"]),errors="coerce")
    contract_finish=pd.to_datetime(get_setup(project_setup,"Contract Completion",bf["finish"]),errors="coerce")
    dash=pd.DataFrame([
        ["Contract Commencement",fmt_date(contract_start),"Contract Completion",fmt_date(contract_finish)],
        ["Forecast Completion",fmt_date(forecast_finish),"Forecast Movement",f"{forecast_move:+.0f} d" if pd.notna(forecast_move) else "—"],
        ["Time Elapsed",f"{elapsed_days:.0f} d","Remaining Time",f"{remaining_days:.0f} d"],
        ["Critical Activities",critical_count,"Negative Float",neg_count],
        ["MHR Planned",f"{mhr_plan:.2f}%" if pd.notna(mhr_plan) else "N/A","MHR Actual",f"{mhr_actual:.2f}%" if pd.notna(mhr_actual) else "N/A"],
        ["Cost Planned",f"{cost_plan:.2f}%" if pd.notna(cost_plan) else "N/A","Cost Actual",f"{cost_actual:.2f}%" if pd.notna(cost_actual) else "N/A"],
    ],columns=["Description","Status","Description 2","Status 2"])
    draw_table(dash,18*mm,y,W*0.45,font=8.5,col_widths=[50*mm,33*mm,50*mm,33*mm])
    summary=pd.DataFrame([
        ["MHR",wk["mhr_wp"],wk["mhr_wa"],wk["mhr_cp"],wk["mhr_ca"],wk["mhr_var"]],
        ["Cost",wk["cost_wp"],wk["cost_wa"],wk["cost_cp"],wk["cost_ca"],wk["cost_var"]],
    ],columns=["Basis","Weekly Plan %","Weekly Actual %","Cumulative Plan %","Cumulative Actual %","Variance %"])
    for cc in summary.columns[1:]: summary[cc]=summary[cc].map(lambda v:"N/A" if pd.isna(v) else f"{v:.2f}%")
    draw_table(summary,W*0.49,y,W*0.48,font=8.2,negative_cols=[5])
    # duration + variance charts
    dd=pd.DataFrame({"Metric":["Original Duration","At Completion Duration","Time Elapsed"],"Days":[duration_days,duration_days+(forecast_move if pd.notna(forecast_move) else 0),elapsed_days]})
    f=px.bar(dd,x="Days",y="Metric",orientation="h",text="Days",title="Project Duration"); f.update_layout(showlegend=False)
    draw_chart(_compare_png(dd["Metric"].tolist(),dd["Days"].tolist(),[0,0,0],"Project Duration"),18*mm,y-64*mm,W*0.45,72*mm)
    var_df=pd.DataFrame({"Basis":["MHR","Cost"],"Plan":[mhr_plan,cost_plan],"Actual":[mhr_actual,cost_actual]})
    vf=go.Figure();vf.add_bar(x=var_df.Basis,y=var_df.Plan,name="Planned",marker_color=BRAND_GREEN);vf.add_bar(x=var_df.Basis,y=var_df.Actual,name="Actual",marker_color=BRAND_YELLOW);vf.update_layout(title="Cumulative Progress")
    draw_chart(_compare_png(var_df.Basis.tolist(),var_df.Plan.fillna(0).tolist(),var_df.Actual.fillna(0).tolist(),"Cumulative Progress"),W*0.49,y-64*mm,W*0.48,72*mm)

    # 5 Progress Summary portrait
    W,H,y=page("2.1 - Progress Summary",False,"Weightage bases are reported separately: Manhours / Units and Cost")
    hist=history.sort_values("Data Date")
    prev=hist.iloc[-2] if len(hist)>1 else None
    cur=hist.iloc[-1] if len(hist)>0 else None
    rows=[]
    for basis in ["MHR","Cost"]:
        if cur is None: rows.append([basis,"N/A"]*1); continue
        last_wp=prev.get(f"{basis} Weekly Plan %",np.nan) if prev is not None else np.nan
        last_wa=prev.get(f"{basis} Weekly Actual %",np.nan) if prev is not None else np.nan
        last_cp=prev.get(f"{basis} Cumulative Plan %",np.nan) if prev is not None else np.nan
        last_ca=prev.get(f"{basis} Cumulative Actual %",np.nan) if prev is not None else np.nan
        rows.append([basis,last_wp,cur.get(f"{basis} Weekly Plan %",np.nan),last_wa,cur.get(f"{basis} Weekly Actual %",np.nan),last_cp,cur.get(f"{basis} Cumulative Plan %",np.nan),last_ca,cur.get(f"{basis} Cumulative Actual %",np.nan),cur.get(f"{basis} Cumulative Actual %",np.nan)-cur.get(f"{basis} Cumulative Plan %",np.nan)])
    ps=pd.DataFrame(rows,columns=["Basis","Weekly Plan Last","Weekly Plan Current","Weekly Actual Last","Weekly Actual Current","Cum Plan Last","Cum Plan Current","Cum Actual Last","Cum Actual Current","Variance"])
    for cc in ps.columns[1:]: ps[cc]=ps[cc].map(lambda v:"N/A" if pd.isna(v) else f"{v:.2f}%")
    draw_table(ps,12*mm,y,W-24*mm,font=6.7,negative_cols=[9])
    # two compact bar charts
    mdf=pd.DataFrame({"Period":["Weekly","Cumulative"],"Planned":[wk["mhr_wp"],wk["mhr_cp"]],"Actual":[wk["mhr_wa"],wk["mhr_ca"]]})
    mf=go.Figure();mf.add_bar(x=mdf.Period,y=mdf.Planned,name="Planned",marker_color=BRAND_GREEN);mf.add_bar(x=mdf.Period,y=mdf.Actual,name="Actual",marker_color=BRAND_YELLOW);mf.update_layout(title="MHR Progress",barmode="group",yaxis_title="%")
    cdf=pd.DataFrame({"Period":["Weekly","Cumulative"],"Planned":[wk["cost_wp"],wk["cost_cp"]],"Actual":[wk["cost_wa"],wk["cost_ca"]]})
    cf=go.Figure();cf.add_bar(x=cdf.Period,y=cdf.Planned,name="Planned",marker_color=BRAND_GREEN);cf.add_bar(x=cdf.Period,y=cdf.Actual,name="Actual",marker_color=BRAND_YELLOW);cf.update_layout(title="Cost Progress",barmode="group",yaxis_title="%")
    draw_chart(_compare_png(mdf.Period.tolist(),mdf.Planned.fillna(0).tolist(),mdf.Actual.fillna(0).tolist(),"MHR Progress"),15*mm,y-65*mm,(W-40*mm)/2,73*mm); draw_chart(_compare_png(cdf.Period.tolist(),cdf.Planned.fillna(0).tolist(),cdf.Actual.fillna(0).tolist(),"Cost Progress"),W/2+5*mm,y-65*mm,(W-40*mm)/2,73*mm)

    # 6 Major trades
    W,H,y=page("2.4 - Major Trades Tracker",True)
    mt=db_sheet(project_db_sheets,"Major Trades")
    if not mt.empty and "Variance %" in mt.columns:
        mt["Variance %"]=pd.to_numeric(mt["Variance %"],errors="coerce").map(lambda v:"" if pd.isna(v) else f"{v*100:.2f}%" if abs(v)<=1 else f"{v:.2f}%")
    draw_table(mt,18*mm,y,W-36*mm,font=8,negative_cols=[4] if not mt.empty and len(mt.columns)>4 else None,max_rows=25)

    # 7 Progress analysis
    W,H,y=page("3 - Progress Analysis",False)
    narrative_sections=["Executive Summary","Key Achievements This Week","Planned Activities Next Week","Key Delays / Constraints","Recovery / Mitigation Measures","Client / Consultant Decisions Required"]
    yy=y
    for sec in narrative_sections:
        c.setFillColor(colors.HexColor(BRAND_GREEN));c.setFont("Helvetica-Bold",11);c.drawString(18*mm,yy,sec)
        yy-=6*mm;c.setFillColor(colors.HexColor("#333333"));c.setFont("Helvetica",8.5)
        text=_db_text(sec); tx=c.beginText(18*mm,yy); tx.setLeading(11)
        for line in str(text).splitlines() or [str(text)]:
            # simple word wrap
            words=line.split(); curline=""
            for word in words:
                if c.stringWidth(curline+" "+word,"Helvetica",8.5) > W-36*mm:
                    tx.textLine(curline);curline=word
                else: curline=(curline+" "+word).strip()
            if curline: tx.textLine(curline)
        c.drawText(tx); yy=tx.getY()-8*mm
        if yy<35*mm: break

    # 8 Labor Force
    W,H,y=page("4 - Labor Force",False)
    actual_mp=db_sheet(project_db_sheets,"Actual Manpower")
    if not actual_mp.empty and "Date" in actual_mp.columns:
        actual_mp["Date"]=pd.to_datetime(actual_mp["Date"],errors="coerce"); actual_mp=actual_mp.dropna(subset=["Date"])
    mpw=manpower_plan_weekly.copy() if manpower_plan_weekly is not None else pd.DataFrame()
    if not mpw.empty:
        mpw["date"]=pd.to_datetime(mpw["date"],errors="coerce")
        if not actual_mp.empty:
            amp=actual_mp.copy(); amp["Total"]=pd.to_numeric(amp.get("Total"),errors="coerce")
            amp=amp.set_index("Date")["Total"].resample("W").mean().reset_index().rename(columns={"Date":"date","Total":"actual"})
            mpw=pd.merge(mpw,amp,on="date",how="left");mpw["actual"]=mpw["actual"].fillna(0)
        else: mpw["actual"]=0
    else:
        mpw=pd.DataFrame(columns=["date","planned","actual"])
    draw_chart(_manpower_png(mpw,"Weekly Planned vs Actual Manpower"),15*mm,y,(W-30*mm),95*mm)
    if not actual_mp.empty:
        show=actual_mp.tail(20).copy(); show["Date"]=show["Date"].dt.strftime("%d-%b-%Y")
        draw_table(show,15*mm,y-105*mm,W-30*mm,font=6.5,max_rows=20)

    # 9 risks
    W,H,y=page("5.1 - AoC & Risk",True)
    risks=db_sheet(project_db_sheets,"Risks & Actions")
    draw_table(risks,15*mm,y,W-30*mm,font=7,max_rows=28)

    # 10 s-curves section cover
    W,H,y=page("6 - Updated S-Curves",True,"Approved planned curves versus actual progress at current data date")
    c.setFillColor(colors.HexColor(BRAND_GREEN));c.setFont("Helvetica-Bold",28);c.drawCentredString(W/2,H/2+15*mm,"S-CURVES / CASH FLOW / MANPOWER")
    c.setFillColor(colors.HexColor("#555555"));c.setFont("Helvetica",13);c.drawCentredString(W/2,H/2,"Monthly and Weekly - Planned vs Actual")

    # 11-16 charts
    cost_m=_curve_resample(cost_curve,"MS"); cost_w=_curve_resample(cost_curve,"W")
    unit_m=_curve_resample(unit_curve,"MS"); unit_w=_curve_resample(unit_curve,"W")
    for title,fig in [
        ("6.1 - Cost S-Curve Monthly",_curve_png(cost_m,"Cost Progress - Monthly")),
        ("6.2 - Cost S-Curve Weekly",_curve_png(cost_w,"Cost Progress - Weekly")),
        ("6.3 - Unit / MHR S-Curve Monthly",_curve_png(unit_m,"MHR / Unit Progress - Monthly")),
        ("6.4 - Unit / MHR S-Curve Weekly",_curve_png(unit_w,"MHR / Unit Progress - Weekly")),
        ("6.1A - Cash Flow Monthly",_cashflow_png(cost_m,"Cash Flow - Monthly Planned vs Actual")),
        ("6.2A - Cash Flow Weekly",_cashflow_png(cost_w,"Cash Flow - Weekly Planned vs Actual")),
    ]:
        W,H,y=page(title,True); draw_chart(fig,18*mm,y,W-36*mm,150*mm)

    # manpower monthly/weekly
    mpm=manpower_plan_monthly.copy() if manpower_plan_monthly is not None else pd.DataFrame()
    mpw2=mpw.copy()
    if not mpm.empty:
        mpm["date"]=pd.to_datetime(mpm["date"],errors="coerce"); mpm["actual"]=0.0
        if not actual_mp.empty:
            ma=actual_mp.set_index("Date")["Total"].resample("MS").mean().reset_index().rename(columns={"Date":"date","Total":"actual"})
            mpm=pd.merge(mpm[["date","planned"]],ma,on="date",how="left");mpm["actual"]=mpm["actual"].fillna(0)
    for title,mpdata in [("6.5 - Manpower Histogram Monthly",mpm),("6.6 - Manpower Histogram Weekly",mpw2)]:
        W,H,y=page(title,True);draw_chart(_manpower_png(mpdata,title),18*mm,y,W-36*mm,150*mm)

    # engineering
    W,H,y=page("Engineering Control - Major Package Wise",True)
    eng=active[active.category.isin(["Prequalification","Shop Drawings","Material Submittals"])].copy()
    eng["Status"]=eng.total_float_days.apply(status_from_float);eng["TF"]=eng.total_float_days.map(lambda v:"—" if pd.isna(v) else f"{v:.0f} d")
    eng["Finish Date"]=eng.finish.map(fmt_date)
    engout=eng[["package","category","Finish Date","TF","Status"]].rename(columns={"package":"Package","category":"Stage"}).sort_values(["Status","Finish Date"])
    draw_table(engout,12*mm,y,W-24*mm,font=6.6,negative_cols=[3],max_rows=44,col_widths=[110*mm,42*mm,30*mm,22*mm,45*mm])

    # procurement
    W,H,y=page("Procurement Action Register",True)
    proc=active[active.category=="Procurement"].copy();proc["Action"]=proc.apply(procurement_action,axis=1,dd=data_date);proc["Finish Date"]=proc.finish.map(fmt_date);proc["TF"]=proc.total_float_days.map(lambda v:"—" if pd.isna(v) else f"{v:.0f} d")
    pout=proc[["task_code","task_name","Finish Date","TF","Action"]].rename(columns={"task_code":"Activity ID","task_name":"Procurement Item"}).sort_values("Action")
    draw_table(pout,12*mm,y,W-24*mm,font=6.5,negative_cols=[3],max_rows=44,col_widths=[32*mm,120*mm,30*mm,22*mm,50*mm])

    # lookahead
    W,H,y=page(f"{lookahead_weeks}-Week Lookahead",True,f"Window: {fmt_date(data_date)} to {fmt_date(data_date+pd.Timedelta(weeks=lookahead_weeks))}")
    la=date_window(active,data_date,lookahead_weeks).copy();la["Start Date"]=la.start.map(fmt_date);la["Finish Date"]=la.finish.map(fmt_date);la["TF"]=la.total_float_days.map(lambda v:"—" if pd.isna(v) else f"{v:.0f} d")
    laout=la[["task_code","task_name","category","Start Date","Finish Date","TF"]].rename(columns={"task_code":"Activity ID","task_name":"Activity Name","category":"Category"})
    draw_table(laout,10*mm,y,W-20*mm,font=6.1,negative_cols=[5],max_rows=48,col_widths=[30*mm,125*mm,35*mm,28*mm,28*mm,18*mm])

    # trackers
    for sheet_name,title in [("Structure Tracker","7.1 - Structure Tracker"),("Finishes Tracker","7.2 - Finishes Tracker")]:
        W,H,y=page(title,True); df=db_sheet(project_db_sheets,sheet_name); draw_table(df,12*mm,y,W-24*mm,font=6.8,max_rows=35)

    # QC
    W,H,y=page("8 - QC Report",False); draw_table(db_sheet(project_db_sheets,"QC"),18*mm,y,W-36*mm,font=8,max_rows=35)
    # RFI
    W,H,y=page("9 - RFI",False); draw_table(db_sheet(project_db_sheets,"RFI"),12*mm,y,W-24*mm,font=6.8,max_rows=38)
    # Commercial
    W,H,y=page("10 - Commercial Status",True); draw_table(db_sheet(project_db_sheets,"Commercial"),18*mm,y,W-36*mm,font=8,max_rows=30)

    # Photos, 4 per portrait page
    photos=progress_photos or []
    if not photos:
        W,H,y=page("11 - Progress Photos",False);c.setFont("Helvetica-Oblique",12);c.setFillColor(colors.HexColor("#666666"));c.drawCentredString(W/2,H/2,"Progress photos not provided")
    else:
        for start in range(0,len(photos),4):
            W,H,y=page("11 - Progress Photos",False,f"Photos {start+1} to {min(start+4,len(photos))}")
            slots=[(18*mm,H-78*mm),(W/2+4*mm,H-78*mm),(18*mm,H/2-12*mm),(W/2+4*mm,H/2-12*mm)]
            pw=W/2-26*mm; ph=118*mm
            for pfile,(x,ytop) in zip(photos[start:start+4],slots):
                try:
                    img=ImageReader(BytesIO(pfile.getvalue())); iw,ih=img.getSize(); ratio=min(pw/iw,ph/ih); ww,hh=iw*ratio,ih*ratio
                    c.drawImage(img,x+(pw-ww)/2,ytop-ph+(ph-hh)/2,width=ww,height=hh,mask='auto')
                    c.setFont("Helvetica",7.5);c.setFillColor(colors.HexColor("#555555"));c.drawString(x,ytop-ph-5*mm,pfile.name[:65])
                except Exception: pass

    c.save();buf.seek(0);return buf.getvalue()

def _ppt_add_logos(slide, prs):
    x=Inches(.35)
    for path,h in [(ALTA_LOGO,.34),(DDDC_LOGO,.34),(DAR_LOGO,.28)]:
        if path.exists():
            try:
                pic=slide.shapes.add_picture(str(path),x,Inches(.18),height=Inches(h));x+=pic.width+Inches(.18)
            except Exception: pass

def _ppt_title(slide,title,subtitle=None):
    _ppt_add_logos(slide,None)
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Inches(.72),Inches(13.333),Inches(.10));bar.fill.solid();bar.fill.fore_color.rgb=RGBColor(7,121,92);bar.line.fill.background()
    tb=slide.shapes.add_textbox(Inches(.45),Inches(.93),Inches(12.2),Inches(.55));p=tb.text_frame.paragraphs[0];p.text=title;p.font.size=Pt(23);p.font.bold=True;p.font.color.rgb=RGBColor(18,18,18)
    if subtitle:
        sb=slide.shapes.add_textbox(Inches(.47),Inches(1.45),Inches(12),Inches(.32));q=sb.text_frame.paragraphs[0];q.text=subtitle;q.font.size=Pt(9);q.font.color.rgb=RGBColor(95,95,95)

def _ppt_table(slide,df,x,y,w,h,max_rows=12,font=10):
    if df is None or df.empty:
        tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(.5));p=tb.text_frame.paragraphs[0];p.text="Data not provided";p.font.size=Pt(12);p.font.italic=True;return
    d=df.head(max_rows).copy(); rows=len(d)+1; cols=len(d.columns)
    table=slide.shapes.add_table(rows,cols,Inches(x),Inches(y),Inches(w),Inches(h)).table
    widths=[w/cols]*cols
    for j,col in enumerate(d.columns):
        table.columns[j].width=Inches(widths[j]);cell=table.cell(0,j);cell.text=str(col);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(7,121,92)
        for p in cell.text_frame.paragraphs:p.font.size=Pt(font);p.font.bold=True;p.font.color.rgb=RGBColor(255,255,255)
    for i,(_,r) in enumerate(d.iterrows(),start=1):
        for j,v in enumerate(r.tolist()):
            cell=table.cell(i,j);cell.text="" if pd.isna(v) else str(v);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(247,247,245) if i%2==0 else RGBColor(255,255,255)
            for p in cell.text_frame.paragraphs:
                p.font.size=Pt(font-1);p.font.color.rgb=RGBColor(214,40,40) if isinstance(v,(int,float,np.integer,np.floating)) and v<0 else RGBColor(18,18,18)

def make_pptx():
    """Executive management deck: visual, decision-oriented, all key programme charts."""
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    project=str(get_setup(project_setup,"Project Name","JUMEIRAH RETAIL DEVELOPMENT"));report_no=get_setup(project_setup,"Report Number","")
    active=curr if curr is not None else base;neg_count=int((active.total_float_days<0).sum());critical_count=int((active.total_float_days<=0).sum())
    history=_history_snapshot();wk=_current_week_summary(history)

    # cover
    s=prs.slides.add_slide(prs.slide_layouts[6])
    if COVER_IMG.exists():s.shapes.add_picture(str(COVER_IMG),0,0,width=prs.slide_width,height=prs.slide_height)
    ov=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(6.4),prs.slide_height);ov.fill.solid();ov.fill.fore_color.rgb=RGBColor(18,18,18);ov.fill.transparency=18;ov.line.fill.background()
    _ppt_add_logos(s,prs)
    tb=s.shapes.add_textbox(Inches(.65),Inches(4.45),Inches(5.5),Inches(1.35));tf=tb.text_frame;p=tf.paragraphs[0];p.text=project;p.font.size=Pt(30);p.font.bold=True;p.font.color.rgb=RGBColor(255,255,255)
    p2=tf.add_paragraph();p2.text=f"WEEKLY PROGRESS REPORT {report_no} | DATA DATE {fmt_date(data_date)}";p2.font.size=Pt(12);p2.font.color.rgb=RGBColor(209,211,212)

    # executive dashboard
    s=prs.slides.add_slide(prs.slide_layouts[6]);_ppt_title(s,"Executive Project Dashboard",f"Data Date {fmt_date(data_date)}")
    metrics=[("Forecast Finish",fmt_date(forecast_finish)),("Movement",f"{forecast_move:+.0f} d" if pd.notna(forecast_move) else "—"),("Critical",str(critical_count)),("Negative Float",str(neg_count)),("MHR Actual",f"{mhr_actual:.2f}%" if pd.notna(mhr_actual) else "N/A"),("Cost Actual",f"{cost_actual:.2f}%" if pd.notna(cost_actual) else "N/A")]
    for i,(k,v) in enumerate(metrics):
        row=i//3;col=i%3;x=.55+col*4.18;y=1.9+row*1.55
        sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(3.75),Inches(1.15));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor(247,247,245);sh.line.color.rgb=RGBColor(209,211,212)
        tf=sh.text_frame;tf.clear();p=tf.paragraphs[0];p.text=k;p.font.size=Pt(10);p.font.color.rgb=RGBColor(90,90,90);q=tf.add_paragraph();q.text=v;q.font.size=Pt(24);q.font.bold=True;q.font.color.rgb=RGBColor(214,40,40) if (k in ["Movement","Negative Float"] and (k!="Movement" or (pd.notna(forecast_move) and forecast_move>0))) else RGBColor(7,121,92)
    summ=pd.DataFrame([["MHR",f"{wk['mhr_wp']:.2f}%" if pd.notna(wk['mhr_wp']) else "N/A",f"{wk['mhr_wa']:.2f}%" if pd.notna(wk['mhr_wa']) else "N/A",f"{wk['mhr_cp']:.2f}%" if pd.notna(wk['mhr_cp']) else "N/A",f"{wk['mhr_ca']:.2f}%" if pd.notna(wk['mhr_ca']) else "N/A",f"{wk['mhr_var']:.2f}%" if pd.notna(wk['mhr_var']) else "N/A"],["Cost",f"{wk['cost_wp']:.2f}%" if pd.notna(wk['cost_wp']) else "N/A",f"{wk['cost_wa']:.2f}%" if pd.notna(wk['cost_wa']) else "N/A",f"{wk['cost_cp']:.2f}%" if pd.notna(wk['cost_cp']) else "N/A",f"{wk['cost_ca']:.2f}%" if pd.notna(wk['cost_ca']) else "N/A",f"{wk['cost_var']:.2f}%" if pd.notna(wk['cost_var']) else "N/A"]],columns=["Basis","Wk Plan","Wk Actual","Cum Plan","Cum Actual","Variance"])
    _ppt_table(s,summ,.55,5.15,12.2,1.45,max_rows=2,font=9)

    def chart_slide(title,fig,subtitle=None):
        sl=prs.slides.add_slide(prs.slide_layouts[6]);_ppt_title(sl,title,subtitle);png=fig if isinstance(fig,(bytes,bytearray)) else _fig_png(fig,1500,760)
        if png:sl.shapes.add_picture(BytesIO(png),Inches(.55),Inches(1.75),width=Inches(12.2),height=Inches(5.25))
        return sl
    cost_m=_curve_resample(cost_curve,"MS");cost_w=_curve_resample(cost_curve,"W");unit_m=_curve_resample(unit_curve,"MS");unit_w=_curve_resample(unit_curve,"W")
    for title,fig in [("Cost S-Curve - Monthly",_curve_png(cost_m,"Cost S-Curve - Monthly")),("Cost S-Curve - Weekly",_curve_png(cost_w,"Cost S-Curve - Weekly")),("MHR / Unit S-Curve - Monthly",_curve_png(unit_m,"MHR / Unit S-Curve - Monthly")),("MHR / Unit S-Curve - Weekly",_curve_png(unit_w,"MHR / Unit S-Curve - Weekly")),("Cash Flow - Planned vs Actual",_cashflow_png(cost_m,"Cash Flow - Planned vs Actual"))]:
        chart_slide(title,fig)

    # manpower
    actual_mp=db_sheet(project_db_sheets,"Actual Manpower");mpw=manpower_plan_weekly.copy() if manpower_plan_weekly is not None else pd.DataFrame()
    if not mpw.empty:
        mpw["date"]=pd.to_datetime(mpw["date"],errors="coerce");mpw["actual"]=0.0
        if not actual_mp.empty and "Date" in actual_mp.columns:
            a=actual_mp.copy();a["Date"]=pd.to_datetime(a["Date"],errors="coerce");a["Total"]=pd.to_numeric(a.get("Total"),errors="coerce");a=a.dropna(subset=["Date"])
            aw=a.set_index("Date")["Total"].resample("W").mean().reset_index().rename(columns={"Date":"date","Total":"actual"});mpw=pd.merge(mpw[["date","planned"]],aw,on="date",how="left").fillna(0)
    chart_slide("Manpower - Planned vs Actual",_manpower_png(mpw,"Manpower - Planned vs Actual"))

    # milestones/schedule
    ms=active[active.duration_days.fillna(-1)==0].copy().sort_values("finish").head(12);ms["Finish"]=ms.finish.map(fmt_date);ms["TF"]=ms.total_float_days.map(lambda v:"—" if pd.isna(v) else f"{v:.0f} d")
    sl=prs.slides.add_slide(prs.slide_layouts[6]);_ppt_title(sl,"Milestones & Schedule Health");_ppt_table(sl,ms[["task_code","task_name","Finish","TF"]].rename(columns={"task_code":"ID","task_name":"Milestone"}),.45,1.75,12.4,4.7,max_rows=12,font=8)

    # engineering
    eng=active[active.category.isin(["Prequalification","Shop Drawings","Material Submittals"])].copy().sort_values(["total_float_days","finish"]);eng["Finish"]=eng.finish.map(fmt_date);eng["TF"]=eng.total_float_days.map(lambda v:"—" if pd.isna(v) else f"{v:.0f} d");eng["Status"]=eng.total_float_days.apply(status_from_float)
    sl=prs.slides.add_slide(prs.slide_layouts[6]);_ppt_title(sl,"Engineering - Major Package Wise","Critical and near-critical PQ / Shop Drawings / Material Submittals");_ppt_table(sl,eng[["package","category","Finish","TF","Status"]].rename(columns={"package":"Package","category":"Stage"}),.35,1.7,12.6,5.25,max_rows=15,font=8)

    # procurement
    proc=active[active.category=="Procurement"].copy();proc["Action"]=proc.apply(procurement_action,axis=1,dd=data_date);proc["Finish"]=proc.finish.map(fmt_date);proc["TF"]=proc.total_float_days.map(lambda v:"—" if pd.isna(v) else f"{v:.0f} d");proc=proc.sort_values(["total_float_days","finish"])
    sl=prs.slides.add_slide(prs.slide_layouts[6]);_ppt_title(sl,"Procurement - Priority Actions","Immediate LPO / overdue / delivery risk");_ppt_table(sl,proc[["task_code","task_name","Finish","TF","Action"]].rename(columns={"task_code":"ID","task_name":"Item"}),.3,1.7,12.7,5.25,max_rows=15,font=8)

    # lookahead
    la=date_window(active,data_date,lookahead_weeks).copy().head(16);la["Start"]=la.start.map(fmt_date);la["Finish"]=la.finish.map(fmt_date);la["TF"]=la.total_float_days.map(lambda v:"—" if pd.isna(v) else f"{v:.0f} d")
    sl=prs.slides.add_slide(prs.slide_layouts[6]);_ppt_title(sl,f"{lookahead_weeks}-Week Lookahead",f"{fmt_date(data_date)} to {fmt_date(data_date+pd.Timedelta(weeks=lookahead_weeks))}");_ppt_table(sl,la[["task_code","task_name","Start","Finish","TF"]].rename(columns={"task_code":"ID","task_name":"Activity"}),.3,1.7,12.7,5.25,max_rows=16,font=7.5)

    # risks + decisions
    risks=db_sheet(project_db_sheets,"Risks & Actions");sl=prs.slides.add_slide(prs.slide_layouts[6]);_ppt_title(sl,"Key Risks, Actions & Decisions Required");_ppt_table(sl,risks,.35,1.7,12.6,3.7,max_rows=8,font=8)
    txt=_db_text("Client / Consultant Decisions Required");tb=sl.shapes.add_textbox(Inches(.55),Inches(5.65),Inches(12.0),Inches(1.0));p=tb.text_frame.paragraphs[0];p.text="DECISIONS REQUIRED";p.font.bold=True;p.font.size=Pt(11);p.font.color.rgb=RGBColor(7,121,92);q=tb.text_frame.add_paragraph();q.text=txt;q.font.size=Pt(10)

    # photos
    if progress_photos:
        for start in range(0,min(len(progress_photos),8),4):
            sl=prs.slides.add_slide(prs.slide_layouts[6]);_ppt_title(sl,"Progress Photos")
            coords=[(.45,1.65),(6.9,1.65),(.45,4.4),(6.9,4.4)]
            for pfile,(x,y) in zip(progress_photos[start:start+4],coords):
                try:sl.shapes.add_picture(BytesIO(pfile.getvalue()),Inches(x),Inches(y),width=Inches(5.9),height=Inches(2.35))
                except Exception:pass

    out=BytesIO();prs.save(out);out.seek(0);return out.getvalue()


with tabs[11]:
    st.markdown('<div class="section-title">Report Center</div>', unsafe_allow_html=True)
    st.write("Generate management outputs from the same data and Data Date shown in the dashboard.")
    c1,c2,c3=st.columns(3)
    with c1:
        try:
            pdf=make_pdf()
            st.download_button("Generate A3 Weekly PDF Report",pdf,
                               f"JRD_WPR_{data_date.strftime('%Y%m%d')}.pdf",
                               "application/pdf",use_container_width=True)
        except Exception as e:
            st.error(f"PDF generation error: {e}")
    with c2:
        try:
            ppt=make_pptx()
            st.download_button("Generate Management Presentation",ppt,
                               f"JRD_Management_Presentation_{data_date.strftime('%Y%m%d')}.pptx",
                               "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                               use_container_width=True)
        except Exception as e:
            st.error(f"PowerPoint generation error: {e}")
    with c3:
        updated_db=make_updated_project_database()
        if updated_db:
            st.download_button("Download Updated Project Database",updated_db,
                               f"JRD_Project_Database_{data_date.strftime('%Y%m%d')}.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               use_container_width=True)

st.markdown("---")
st.caption("JRD Project Controls Hub · v0.5 WPR MASTER REPORTING · DDDC branded · Negative float shown in red as delay indication")
